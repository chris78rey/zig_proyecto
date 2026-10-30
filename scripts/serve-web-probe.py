#!/usr/bin/env python3
"""Servidor local para probar Oracle → Parquet desde el navegador."""
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import json
from io import BytesIO
import hashlib
import mimetypes
import os
import re
import secrets
import sqlite3
import shutil
import subprocess
import tempfile
import threading
import time
import uuid
from urllib.parse import parse_qs, unquote, urlsplit

import oracledb

ROOT = Path(__file__).resolve().parents[1]
WEB = ROOT / "web"
PROBE = ROOT / "scripts" / "oracle-read-probe.sh"
SIS_ITEMS_EXCEL = ROOT / "outputs" / "SIS_ITEMS_TIPO_I.xlsx"
ITEM_EXPORT_ROOT = ROOT / "output" / "items-exports"
ITEM_REVIEW_DB = ROOT / "output" / "items-review" / "review.sqlite3"
ITEM_PRIMARY_KEY_FIELDS = ("TIPO", "SBS_SCC_CODIGO", "SBS_CODIGO", "CODIGO")
HOST, PORT = "127.0.0.1", 8765
MAX_OUTPUT = 8192
DOWNLOAD_TTL = 15 * 60
EXCEL_NUMERIC_COLUMNS = {
    "STOCK1", "STOCK2", "STOCK_COMPROMETIDO", "STOCK_MINIMO", "COSTO_PROMEDIO",
    "PRECIO_REPOSICION", "PRECIO_DE_VENTA", "IVA", "PORCENTAJE_ADMINISTRACION",
    "PRC_MAX_DESCUENTO", "DIAS_STOCK", "FACTOR_DIVISION", "PRC_RECARGO_COMPRA",
    "STOCK_MAXIMO", "DIAS_STOCK_MAXIMO", "CONSUMO_PROMEDIO", "PRESENTACION",
}
IDENTIFIER = re.compile(r"^[A-Za-z][A-Za-z0-9_$#]{0,29}$")
EXPORT_RE = re.compile(
    r"Exportación real correcta: ([A-Z0-9_$#]+)\.([A-Z0-9_$#]+), filas=(\d+), "
    r"archivo=([^,]+), bytes=(\d+), limite=([^,]+), tiempo_ejecucion=([\d.]+) ms\."
)
VERIFY_RE = re.compile(r"Verificación PyArrow correcta: filas=(\d+), columnas=(\d+), formato=Parquet\.")


def load_local_env():
    env_file = ROOT / ".env"
    if not env_file.is_file():
        return
    for line in env_file.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.removeprefix("export ").strip()
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
            value = value[1:-1]
        os.environ.setdefault(key, value)


load_local_env()
oracledb.init_oracle_client(lib_dir=os.environ.get("ORACLE_CLIENT_LIB_DIR", "/usr/lib/oracle/19.10/client64/lib"))
downloads = {}
downloads_lock = threading.Lock()
full_jobs = {}
full_jobs_lock = threading.Lock()
full_export_running = False
items_jobs = {}
items_jobs_lock = threading.Lock()
items_export_running = False
items_compare_jobs = {}
items_compare_running = False
items_review_lock = threading.Lock()
request_context = threading.local()
sessions = {}
sessions_lock = threading.Lock()
SESSION_TTL = 8 * 60 * 60
COMPARISON_RUNNER_USER = os.environ.get("ITEMS_COMPARE_RUNNER_USER", "christian_ruiz").strip().casefold()


def init_review_db():
    ITEM_REVIEW_DB.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        db.execute("""CREATE TABLE IF NOT EXISTS item_decisions (
            pair_id TEXT PRIMARY KEY, decision TEXT NOT NULL, keep_side TEXT,
            canonical_pk TEXT, saved_by TEXT NOT NULL, saved_at TEXT NOT NULL
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS comparison_runs (
            id TEXT PRIMARY KEY, tipo TEXT NOT NULL, threshold INTEGER NOT NULL,
            auto_same_threshold INTEGER NOT NULL DEFAULT 94,
            state TEXT NOT NULL, created_by TEXT NOT NULL, created_at TEXT NOT NULL,
            completed_at TEXT, rows_processed INTEGER NOT NULL DEFAULT 0,
            candidate_count INTEGER NOT NULL DEFAULT 0, retained_count INTEGER NOT NULL DEFAULT 0,
            error TEXT, deleted_by TEXT, deleted_at TEXT
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS comparison_candidates (
            run_id TEXT NOT NULL, ordinal INTEGER NOT NULL, pair_id TEXT NOT NULL,
            score REAL NOT NULL, item_a TEXT NOT NULL, item_b TEXT NOT NULL,
            PRIMARY KEY(run_id, ordinal), UNIQUE(run_id, pair_id)
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS item_decision_audit (
            id INTEGER PRIMARY KEY AUTOINCREMENT, run_id TEXT NOT NULL,
            pair_id TEXT NOT NULL, decision TEXT NOT NULL, keep_side TEXT,
            canonical_pk TEXT, modified_by TEXT NOT NULL, modified_at TEXT NOT NULL
        )""")
        run_columns = {row[1] for row in db.execute("PRAGMA table_info(comparison_runs)")}
        if "deleted_by" not in run_columns:
            db.execute("ALTER TABLE comparison_runs ADD COLUMN deleted_by TEXT")
        if "deleted_at" not in run_columns:
            db.execute("ALTER TABLE comparison_runs ADD COLUMN deleted_at TEXT")
        if "auto_same_threshold" not in run_columns:
            db.execute("ALTER TABLE comparison_runs ADD COLUMN auto_same_threshold INTEGER NOT NULL DEFAULT 94")
        db.execute("""CREATE TABLE IF NOT EXISTS comparison_run_audit (
            id INTEGER PRIMARY KEY AUTOINCREMENT, run_id TEXT NOT NULL,
            action TEXT NOT NULL, modified_by TEXT NOT NULL, modified_at TEXT NOT NULL
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS item_assignments (
            run_id TEXT NOT NULL, pair_id TEXT NOT NULL, assigned_to TEXT NOT NULL,
            assigned_by TEXT NOT NULL, assigned_at TEXT NOT NULL,
            PRIMARY KEY(run_id,pair_id)
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS item_assignment_audit (
            id INTEGER PRIMARY KEY AUTOINCREMENT, run_id TEXT NOT NULL,
            pair_id TEXT NOT NULL, assigned_to TEXT NOT NULL,
            assigned_by TEXT NOT NULL, assigned_at TEXT NOT NULL
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS item_group_assignments (
            run_id TEXT NOT NULL, group_id TEXT NOT NULL, assigned_to TEXT NOT NULL,
            assigned_by TEXT NOT NULL, assigned_at TEXT NOT NULL,
            PRIMARY KEY(run_id,group_id)
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS item_group_assignment_audit (
            id INTEGER PRIMARY KEY AUTOINCREMENT, run_id TEXT NOT NULL, group_id TEXT NOT NULL,
            assigned_to TEXT NOT NULL, assigned_by TEXT NOT NULL, assigned_at TEXT NOT NULL
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS item_group_decisions (
            run_id TEXT NOT NULL, group_id TEXT NOT NULL, decision TEXT NOT NULL,
            canonical_pk TEXT, members_json TEXT NOT NULL,
            saved_by TEXT NOT NULL, saved_at TEXT NOT NULL,
            PRIMARY KEY(run_id,group_id)
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS item_group_decision_audit (
            id INTEGER PRIMARY KEY AUTOINCREMENT, run_id TEXT NOT NULL,
            group_id TEXT NOT NULL, decision TEXT NOT NULL, canonical_pk TEXT,
            modified_by TEXT NOT NULL, modified_at TEXT NOT NULL
        )""")
        for table in ("item_group_decisions", "item_group_decision_audit"):
            columns = {row[1] for row in db.execute(f"PRAGMA table_info({table})")}
            if "source_group_id" not in columns:
                db.execute(f"ALTER TABLE {table} ADD COLUMN source_group_id TEXT")
            db.execute(f"UPDATE {table} SET source_group_id=group_id WHERE source_group_id IS NULL")
        audit_columns = {row[1] for row in db.execute("PRAGMA table_info(item_group_decision_audit)")}
        if "members_json" not in audit_columns:
            db.execute("ALTER TABLE item_group_decision_audit ADD COLUMN members_json TEXT")
        db.execute("""UPDATE item_group_decision_audit SET members_json=(
            SELECT members_json FROM item_group_decisions
            WHERE item_group_decisions.run_id=item_group_decision_audit.run_id
              AND item_group_decisions.group_id=item_group_decision_audit.group_id
        ) WHERE members_json IS NULL""")
        db.execute("CREATE INDEX IF NOT EXISTS idx_group_audit_run ON item_group_decision_audit(run_id, group_id, id)")
        db.execute("CREATE INDEX IF NOT EXISTS idx_comparison_runs_created ON comparison_runs(created_at DESC)")
        db.execute("CREATE INDEX IF NOT EXISTS idx_decision_audit_pair ON item_decision_audit(pair_id, id DESC)")
        db.execute("UPDATE comparison_runs SET state='interrupted', error='El servidor se reinició antes de completar la comparación.' WHERE state IN ('queued','reading','comparing')")


init_review_db()


def oracle_connection():
    credentials = getattr(request_context, "credentials", None)
    user = credentials[0] if credentials else os.environ.get("ORACLE_USER", "")
    password = credentials[1] if credentials else os.environ.get("ORACLE_PASSWORD", "")
    host = os.environ.get("ORACLE_HOST", "")
    sid = os.environ.get("ORACLE_SID", "")
    if not user or not password:
        raise RuntimeError("missing Oracle credentials")
    if not host or not sid:
        raise RuntimeError("missing Oracle host or SID configuration")
    port = int(os.environ.get("ORACLE_PORT", "1521"))
    dsn = oracledb.makedsn(host, port, sid=sid)
    return oracledb.connect(user=user, password=password, dsn=dsn)


def valid_identifier(value):
    return bool(IDENTIFIER.fullmatch(value))


def estimate_table(schema, table):
    with oracle_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT num_rows, avg_row_len, last_analyzed FROM all_tables "
                "WHERE owner = :owner AND table_name = :table_name",
                owner=schema, table_name=table,
            )
            row = cursor.fetchone()
    if row is None or row[0] is None or row[1] is None:
        raise ValueError("No hay estadísticas de tamaño para esa tabla.")
    estimated_rows, avg_row_len, last_analyzed = row
    estimated_bytes = int(estimated_rows * avg_row_len)
    return {
        "estimated_rows": int(estimated_rows),
        "average_row_bytes": int(avg_row_len),
        "estimated_raw_bytes": estimated_bytes,
        "last_analyzed": last_analyzed.isoformat() if last_analyzed else None,
        "free_bytes": shutil.disk_usage(ROOT).free,
    }


def run_full_export(job_id, schema, table, output_dir, estimate, credentials):
    global full_export_running
    started = time.monotonic()
    env = {
        **os.environ,
        "ORACLE_SCHEMA": schema,
        "ORACLE_TABLE": table,
        "ORACLE_MAX_ROWS": "all",
        "ORACLE_FETCH_ARRAY_SIZE": "1000",
        "EXPORT_PROGRESS": "1",
        "PARQUET_OUTPUT_DIR": str(output_dir),
        "ORACLE_USER": credentials[0], "ORACLE_PASSWORD": credentials[1],
    }
    with full_jobs_lock:
        full_jobs[job_id]["state"] = "connecting"
        full_jobs[job_id]["started_at"] = started
    try:
        process = subprocess.Popen(
            ["bash", str(PROBE)], cwd=ROOT, env=env,
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
            text=True, bufsize=1,
        )
        with full_jobs_lock:
            full_jobs[job_id]["process"] = process
            full_jobs[job_id]["state"] = "reading"
        output_tail = ""
        for line in process.stdout:
            output_tail = (output_tail + line)[-MAX_OUTPUT:]
            progress = re.search(r"EXPORT_PROGRESS rows=(\d+) batch=(\d+) state=([a-z_]+)", line)
            if progress:
                with full_jobs_lock:
                    job = full_jobs[job_id]
                    job["rows_processed"] = int(progress.group(1))
                    job["current_batch"] = int(progress.group(2))
                    job["state"] = progress.group(3)
                    job["elapsed_ms"] = round((time.monotonic() - started) * 1000)
        return_code = process.wait()
        exported = EXPORT_RE.search(output_tail)
        verified = VERIFY_RE.search(output_tail)
        parquet_file = output_dir / "part-00001.parquet"
        if return_code != 0 or not exported or not verified or not parquet_file.is_file():
            shutil.rmtree(output_dir, ignore_errors=True)
            with full_jobs_lock:
                full_jobs[job_id].update({"state": "failed", "error": "Falló la exportación o la verificación Parquet."})
            return
        with full_jobs_lock:
            full_jobs[job_id].update({
                "state": "completed",
                "rows_processed": int(exported.group(3)),
                "columns": int(verified.group(2)),
                "bytes_written": int(exported.group(5)),
                "elapsed_ms": round(float(exported.group(7))),
                "output_file": str(parquet_file.relative_to(ROOT)),
                "download_url": f"/api/v1/full-exports/{job_id}/download",
                "statistics": estimate,
            })
    except Exception:
        shutil.rmtree(output_dir, ignore_errors=True)
        with full_jobs_lock:
            full_jobs[job_id].update({"state": "failed", "error": "No se pudo completar la exportación."})
    finally:
        with full_jobs_lock:
            full_export_running = False


def item_type_counts():
    with oracle_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT TIPO, COUNT(*) FROM SIS.ITEMS WHERE TIPO IS NOT NULL "
                "GROUP BY TIPO ORDER BY TIPO"
            )
            return [{"tipo": str(tipo), "rows": int(rows)} for tipo, rows in cursor]


def oracle_reviewer_users():
    with oracle_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute("SELECT username FROM all_users ORDER BY username")
            return [str(row[0]) for row in cursor if row[0]]


def item_type_row_count(tipo):
    with oracle_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) FROM SIS.ITEMS WHERE TIPO = :tipo", tipo=tipo)
            return int(cursor.fetchone()[0])


def convert_items_parquet_to_excel(parquet_path, excel_path, tipo, row_count):
    import pyarrow.parquet as pq
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from datetime import datetime, timezone

    parquet = pq.ParquetFile(parquet_path)
    if parquet.metadata.num_rows > 1_048_573:
        raise ValueError("El tipo seleccionado supera el límite de filas de una hoja Excel.")
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Items")
    sheet.append([f"Oracle SIS.ITEMS | TIPO={tipo} | Registros: {row_count}"])
    sheet.append([])
    headers = parquet.schema_arrow.names
    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(sheet, value=header)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        header_cells.append(cell)
    sheet.append(header_cells)
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:{excel_column(len(headers))}{row_count + 3}"
    for batch in parquet.iter_batches(batch_size=1000):
        for row in batch.to_pylist():
            cells = []
            for header, value in row.items():
                if isinstance(value, datetime) and value.tzinfo is not None:
                    value = value.astimezone(timezone.utc).replace(tzinfo=None)
                if header in EXCEL_NUMERIC_COLUMNS and isinstance(value, str):
                    try:
                        from decimal import Decimal
                        decimal_value = Decimal(value)
                        digits = len(decimal_value.as_tuple().digits)
                        if decimal_value == decimal_value.to_integral_value() and abs(decimal_value) <= 9_007_199_254_740_991:
                            value = int(decimal_value)
                        elif abs(decimal_value) < Decimal("1e12") and digits <= 15:
                            value = float(decimal_value)
                    except Exception:
                        pass
                if isinstance(value, str):
                    cell = WriteOnlyCell(sheet, value=value)
                    cell.data_type = "s"
                    cells.append(cell)
                else:
                    cells.append(value)
            sheet.append(cells)
    sheet.column_dimensions["E"].width = 52
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["AB"].width = 22
    workbook.save(excel_path)


def excel_column(index):
    value = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        value = chr(65 + remainder) + value
    return value


def normalize_item_description(value):
    import unicodedata
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char)).upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def compare_item_records(records, threshold, limit=10_000):
    from rapidfuzz import fuzz, process

    usable = [(row, normalize_item_description(row.get("DESCRIPCION"))) for row in records]
    usable = [(row, text) for row, text in usable if text]
    choices = [text for _, text in usable]
    sorted_choices = [" ".join(sorted(text.split())) for text in choices]
    candidate_pairs = set()
    for index, (text, sorted_text) in enumerate(zip(choices, sorted_choices)):
        for query, pool in ((text, choices), (sorted_text, sorted_choices)):
            for _, _, other_index in process.extract(
                query, pool, scorer=fuzz.ratio, processor=None,
                score_cutoff=threshold, limit=8,
            ):
                if other_index != index:
                    candidate_pairs.add((min(index, other_index), max(index, other_index)))

    scored = []
    for left_index, right_index in candidate_pairs:
        left, left_text = usable[left_index]
        right, right_text = usable[right_index]
        character_score = fuzz.ratio(left_text, right_text)
        token_score = fuzz.ratio(sorted_choices[left_index], sorted_choices[right_index])
        score = round(character_score * 0.65 + token_score * 0.35, 1)
        if score < threshold:
            continue
        left_pk = tuple(str(left.get(field) or "") for field in ITEM_PRIMARY_KEY_FIELDS)
        right_pk = tuple(str(right.get(field) or "") for field in ITEM_PRIMARY_KEY_FIELDS)
        if right_pk < left_pk:
            left, right = right, left
            left_pk, right_pk = right_pk, left_pk
        fields = (*ITEM_PRIMARY_KEY_FIELDS, "DESCRIPCION", "UNIDAD", "ESTADO", "STOCK1", "STOCK2", "STOCK_COMPROMETIDO", "COSTO_PROMEDIO", "PRECIO_DE_VENTA")
        scored.append({
            "score": score,
            "a": {field: left.get(field) for field in fields},
            "b": {field: right.get(field) for field in fields},
        })
    scored.sort(key=lambda pair: (-pair["score"], tuple(str(pair["a"].get(f) or "") for f in ITEM_PRIMARY_KEY_FIELDS), tuple(str(pair["b"].get(f) or "") for f in ITEM_PRIMARY_KEY_FIELDS)))
    return scored[:limit], len(scored)


def review_pair_id(pair):
    keys = [tuple(str(pair[side].get(field) or "") for field in ITEM_PRIMARY_KEY_FIELDS) for side in ("a", "b")]
    keys.sort()
    return json.dumps(keys, ensure_ascii=False, separators=(",", ":"))


def read_item_decisions():
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        rows = db.execute("SELECT pair_id, decision, keep_side, canonical_pk, saved_by, saved_at FROM item_decisions")
        return {key: {"decision": decision, "keep": keep_side,
                      "canonical_pk": json.loads(canonical_pk) if canonical_pk else None,
                      "saved_by": saved_by, "saved_at": saved_at}
                for key, decision, keep_side, canonical_pk, saved_by, saved_at in rows}


def attach_item_decisions(pairs):
    refresh_pair_decisions(pairs)
    return pairs


def refresh_pair_decisions(pairs):
    if not pairs:
        return
    ids = [review_pair_id(pair) for pair in pairs]
    saved = {}
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        for offset in range(0, len(ids), 900):
            batch = ids[offset:offset + 900]
            slots = ",".join("?" for _ in batch)
            rows = db.execute(
                f"SELECT pair_id, decision, keep_side, canonical_pk, saved_by, saved_at FROM item_decisions WHERE pair_id IN ({slots})",
                batch,
            )
            for key, decision, keep, canonical, user, saved_at in rows:
                saved[key] = {"decision": decision, "keep": keep,
                              "canonical_pk": json.loads(canonical) if canonical else None,
                              "saved_by": user, "saved_at": saved_at}
    for pair in pairs:
        state = saved.get(review_pair_id(pair))
        pair["decision"] = state.get("decision") if state else None
        pair["keep"] = state.get("keep") if state else None
        pair["canonical_pk"] = state.get("canonical_pk") if state else None
        pair["saved_by"] = state.get("saved_by") if state else None
        pair["saved_at"] = state.get("saved_at") if state else None
        pair["price_difference"] = calculate_price_difference(pair)


def calculate_price_difference(pair):
    from decimal import Decimal, InvalidOperation

    def amount(value):
        raw = str(value or "").strip()
        if not raw:
            return None
        if "," in raw and "." not in raw:
            raw = raw.replace(",", ".")
        else:
            raw = raw.replace(",", "")
        try:
            return Decimal(raw)
        except InvalidOperation:
            return None

    price_a = amount(pair["a"].get("PRECIO_DE_VENTA"))
    price_b = amount(pair["b"].get("PRECIO_DE_VENTA"))
    if price_a is None or price_b is None:
        return None
    delta = price_b - price_a
    if price_a == 0:
        percentage = None
        high_difference = price_b != 0
    else:
        percentage = delta / price_a * Decimal("100")
        high_difference = abs(percentage) > Decimal("15")
    return {
        "field": "PRECIO_DE_VENTA", "a": str(price_a), "b": str(price_b),
        "delta": str(delta), "percentage_vs_a": str(percentage) if percentage is not None else None,
        "high_difference": high_difference,
    }


def refresh_pair_assignments(run_id, pairs):
    if not pairs:
        return
    ids = [review_pair_id(pair) for pair in pairs]
    assignments = {}
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        for offset in range(0, len(ids), 900):
            batch = ids[offset:offset + 900]
            slots = ",".join("?" for _ in batch)
            rows = db.execute(
                f"SELECT pair_id,assigned_to,assigned_by,assigned_at FROM item_assignments WHERE run_id=? AND pair_id IN ({slots})",
                [run_id, *batch],
            )
            for pair_id, assigned_to, assigned_by, assigned_at in rows:
                assignments[pair_id] = {"assigned_to": assigned_to, "assigned_by": assigned_by, "assigned_at": assigned_at}
    for pair in pairs:
        pair["assignment"] = assignments.get(review_pair_id(pair))


def item_pk_tuple(item):
    return tuple("" if item.get(field) is None else str(item.get(field)) for field in ITEM_PRIMARY_KEY_FIELDS)


def build_comparison_groups(run_id, pairs, min_score):
    """Build transitive candidate groups from edges meeting the run's suggestion threshold."""
    parent, items = {}, {}

    def find(key):
        parent.setdefault(key, key)
        if parent[key] != key:
            parent[key] = find(parent[key])
        return parent[key]

    def union(left, right):
        root_left, root_right = find(left), find(right)
        if root_left != root_right:
            parent[root_right] = root_left

    for pair in pairs:
        a_key, b_key = item_pk_tuple(pair["a"]), item_pk_tuple(pair["b"])
        items.setdefault(a_key, pair["a"])
        items.setdefault(b_key, pair["b"])
        if float(pair.get("score") or 0) >= min_score:
            union(a_key, b_key)

    components = {}
    for key in list(parent):
        components.setdefault(find(key), set()).add(key)
    key_to_component = {}
    result = []
    for members in components.values():
        if len(members) < 2:
            continue
        sorted_keys = sorted(members)
        digest_input = run_id + json.dumps(sorted_keys, ensure_ascii=False, separators=(",", ":"))
        group_id = hashlib.sha256(digest_input.encode("utf-8")).hexdigest()[:24]
        for key in members:
            key_to_component[key] = group_id
        result.append({
            "id": group_id,
            "member_keys": sorted_keys,
            "members": [{"pk": {field: items[key].get(field) for field in ITEM_PRIMARY_KEY_FIELDS}, "item": items[key]}
                        for key in sorted_keys],
            "links": [],
            "max_score": 0,
            "min_score": 100,
        })

    groups_by_id = {group["id"]: group for group in result}
    for pair in pairs:
        a_key, b_key = item_pk_tuple(pair["a"]), item_pk_tuple(pair["b"])
        group_id = key_to_component.get(a_key)
        if group_id is None or key_to_component.get(b_key) != group_id:
            continue
        group = groups_by_id[group_id]
        score = float(pair.get("score") or 0)
        group["max_score"] = max(group["max_score"], score)
        group["min_score"] = min(group["min_score"], score)
        group["links"].append({
            "pair_id": review_pair_id(pair),
            "score": score,
            "a_pk": {field: pair["a"].get(field) for field in ITEM_PRIMARY_KEY_FIELDS},
            "b_pk": {field: pair["b"].get(field) for field in ITEM_PRIMARY_KEY_FIELDS},
            "a_description": pair["a"].get("DESCRIPCION"),
            "b_description": pair["b"].get("DESCRIPCION"),
            "decision": pair.get("decision"),
            "saved_by": pair.get("saved_by"),
            "assignment": pair.get("assignment"),
        })
    for group in result:
        group["links"].sort(key=lambda link: (-link["score"], item_pk_tuple(link["a_pk"]), item_pk_tuple(link["b_pk"])))
    result.sort(key=lambda group: (-group["max_score"], -len(group["members"]), group["id"]))
    return result


def attach_group_decisions(run_id, groups):
    if not groups:
        return
    ids = [group["id"] for group in groups]
    states, histories = {}, {group_id: [] for group_id in ids}
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        for offset in range(0, len(ids), 900):
            batch = ids[offset:offset + 900]
            slots = ",".join("?" for _ in batch)
            for row in db.execute(
                f"SELECT group_id,source_group_id,decision,canonical_pk,saved_by,saved_at,members_json FROM item_group_decisions WHERE run_id=? AND source_group_id IN ({slots})",
                [run_id, *batch],
            ):
                source = row[1] or row[0]
                states.setdefault(source, []).append({"id": row[0], "decision": row[2], "canonical_pk": json.loads(row[3]) if row[3] else None,
                                  "saved_by": row[4], "saved_at": row[5], "members_json": row[6]})
            for row in db.execute(
                f"SELECT group_id,source_group_id,decision,canonical_pk,modified_by,modified_at,members_json FROM item_group_decision_audit WHERE run_id=? AND source_group_id IN ({slots}) ORDER BY id",
                [run_id, *batch],
            ):
                histories.setdefault(row[1] or row[0], []).append({
                    "decision": row[2], "canonical_pk": json.loads(row[3]) if row[3] else None,
                    "modified_by": row[4], "modified_at": row[5],
                    "member_pks": json.loads(row[6]) if row[6] else [],
                })
    for group in groups:
        # Only explicit same-item or skip decisions close members. Legacy manual
        # decisions were a routing suggestion, so make those members reviewable again.
        group["subgroups"] = [state for state in states.get(group["id"], []) if state["decision"] == "unify"]
        skipped = [state for state in states.get(group["id"], []) if state["decision"] == "skip"]
        group["skipped_members"] = [
            {"pk": pk, "saved_by": state["saved_by"], "saved_at": state["saved_at"]}
            for state in skipped for pk in json.loads(state["members_json"])
        ]
        closed_keys = {
            item_pk_tuple(pk)
            for state in group["subgroups"] + skipped
            for pk in json.loads(state["members_json"])
        }
        group["pending_keys"] = [member["pk"] for member in group["members"]
                                 if item_pk_tuple(member["pk"]) not in closed_keys]
        group["review"] = None
        group["audit_history"] = histories.get(group["id"], [])


def group_merge_plan_workbook(run_id):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    run = load_comparison_run(run_id)
    if run is None or run.get("state") != "completed":
        return None
    refresh_pair_assignments(run_id, run["results"])
    groups = build_comparison_groups(run_id, run["results"], run.get("auto_same_threshold", 94))
    attach_group_decisions(run_id, groups)
    mapping_headers = [
        "TIPO_ORIGEN", "SBS_SCC_CODIGO_ORIGEN", "SBS_CODIGO_ORIGEN", "CODIGO_ORIGEN",
        "DESCRIPCION_ORIGEN", "TIPO_MAESTRO", "SBS_SCC_CODIGO_MAESTRO",
        "SBS_CODIGO_MAESTRO", "CODIGO_MAESTRO", "DESCRIPCION_MAESTRO", "ACCION",
        "ID_CORRIDA", "ID_GRUPO", "ID_SUBGRUPO", "USUARIO_ASIGNADO", "DECIDIDO_POR",
        "FECHA_DECISION", "PUNTAJE_SUBGRUPO",
    ]
    pending_headers = [
        "ARTICULOS_CONFIRMADOS", "ARTICULOS_DEL_SUBGRUPO", "USUARIO_ASIGNADO",
        "DECIDIDO_POR", "FECHA_DECISION", "ID_CORRIDA", "ID_GRUPO", "ID_SUBGRUPO",
    ]
    mapping_rows, pending_rows = [], []
    for group in groups:
        assignees = [(link.get("assignment") or {}).get("assigned_to") for link in group["links"]]
        assigned_to = assignees[0] if assignees and all(
            user and user.casefold() == assignees[0].casefold() for user in assignees
        ) else ""
        for review in group.get("subgroups", []):
            if review.get("decision") != "unify":
                continue
            selected_keys = {item_pk_tuple(pk) for pk in json.loads(review["members_json"])}
            selected_members = [member for member in group["members"] if item_pk_tuple(member["pk"]) in selected_keys]
            canonical_tuple = item_pk_tuple(review["canonical_pk"]) if isinstance(review.get("canonical_pk"), dict) else None
            canonical_member = next((member for member in selected_members if item_pk_tuple(member["pk"]) == canonical_tuple), None)
            if canonical_member is None:
                descriptions = []
                for member in selected_members:
                    key_text = " / ".join(f"{field}={member['pk'].get(field) or ''}" for field in ITEM_PRIMARY_KEY_FIELDS)
                    descriptions.append(f"{member['item'].get('DESCRIPCION') or 'Sin descripción'} [{key_text}]")
                member_summary = " · ".join(descriptions)
                pending_rows.append([
                    len(selected_members), member_summary, assigned_to, review.get("saved_by"),
                    review.get("saved_at"), run_id, group["id"], review["id"],
                ])
                continue
            canonical = canonical_member["item"]
            subgroup_scores = [
                float(link["score"]) for link in group["links"]
                if item_pk_tuple(link["a_pk"]) in selected_keys and item_pk_tuple(link["b_pk"]) in selected_keys
            ]
            subgroup_score = max(subgroup_scores, default=0)
            for member in selected_members:
                if item_pk_tuple(member["pk"]) == canonical_tuple:
                    continue
                source = member["item"]
                values = list(source.get(field) for field in ITEM_PRIMARY_KEY_FIELDS)
                values.append(source.get("DESCRIPCION"))
                values.extend(canonical.get(field) for field in ITEM_PRIMARY_KEY_FIELDS)
                values.extend([canonical.get("DESCRIPCION"), "EMPATAR CON MAESTRO", run_id, group["id"], review["id"],
                               assigned_to, review.get("saved_by"), review.get("saved_at"), subgroup_score])
                mapping_rows.append(values)

    workbook = Workbook()
    mapping_sheet = workbook.active
    mapping_sheet.title = "Empates a realizar"
    pending_sheet = workbook.create_sheet("Pendientes de maestro")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    title_fill = PatternFill("solid", fgColor="176B55")
    for sheet, title, instruction, headers, rows in (
        (mapping_sheet, "PLAN DE EMPATES · REQUIERE REVISIÓN TÉCNICA",
         "Cada fila indica qué artículo origen se debe asociar a la clave del maestro. Este reporte no ejecuta cambios en Oracle.", mapping_headers, mapping_rows),
        (pending_sheet, "SUBGRUPOS PENDIENTES DE ARTÍCULO MAESTRO",
         "Confirma el artículo que se conservará en cada subgrupo antes de preparar todos los empates.", pending_headers, pending_rows),
    ):
        sheet.sheet_view.showGridLines = False
        sheet.append([title])
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        sheet.append([instruction])
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        sheet.append(headers)
        sheet.row_dimensions[1].height = 28
        sheet.row_dimensions[2].height = 34
        sheet.row_dimensions[3].height = 42
        sheet[1][0].font = Font(bold=True, size=15, color="FFFFFF")
        sheet[1][0].fill = title_fill
        sheet[2][0].font = Font(italic=True, color="365854")
        sheet[2][0].alignment = Alignment(wrap_text=True, vertical="center")
        for cell in sheet[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        if sheet is mapping_sheet:
            for column_index in range(6, 11):
                sheet.cell(row=3, column=column_index).fill = PatternFill("solid", fgColor="176B55")
        for row_values in rows:
            row_index = sheet.max_row + 1
            for column_index, value in enumerate(row_values, 1):
                cell = sheet.cell(row=row_index, column=column_index, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=column_index in (5, 10) if sheet is mapping_sheet else column_index == 2)
                if sheet is mapping_sheet and 6 <= column_index <= 10:
                    cell.fill = PatternFill("solid", fgColor="EAF6EF")
                text_columns = (1, 2, 3, 4, 6, 7, 8, 9, 12, 13, 14) if sheet is mapping_sheet else (6, 7, 8)
                if column_index in text_columns and value is not None:
                    cell.value = str(value)
                    cell.number_format = "@"
                if sheet is mapping_sheet and column_index == 18:
                    cell.number_format = '0.0"%"'
            long_text = max((len(str(value)) for value in row_values if value is not None), default=0)
            sheet.row_dimensions[row_index].height = min(120, max(34, 15 * ((long_text + 89) // 90)))
        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = f"A3:{excel_column(len(headers))}{max(3, sheet.max_row)}"
    for column, width in {"A": 14, "B": 27, "C": 23, "D": 17, "E": 42, "F": 14, "G": 27, "H": 23, "I": 17, "J": 42, "K": 24, "L": 35, "M": 27, "N": 27, "O": 24, "P": 24, "Q": 23, "R": 18}.items():
        mapping_sheet.column_dimensions[column].width = width
    for column, width in {"A": 35, "B": 27, "C": 27, "D": 20, "E": 90, "F": 24, "G": 24, "H": 23}.items():
        pending_sheet.column_dimensions[column].width = width
    mapping_sheet.page_setup.orientation = "landscape"
    pending_sheet.page_setup.orientation = "landscape"
    pending_sheet.page_setup.fitToWidth = 1
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def assign_run_pairs(run_id, pairs, assigned_to, assigned_by):
    now = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    records = [(run_id, review_pair_id(pair), assigned_to, assigned_by, now) for pair in pairs]
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        db.executemany("""INSERT INTO item_assignments (run_id,pair_id,assigned_to,assigned_by,assigned_at)
            VALUES (?,?,?,?,?) ON CONFLICT(run_id,pair_id) DO UPDATE SET
            assigned_to=excluded.assigned_to,assigned_by=excluded.assigned_by,assigned_at=excluded.assigned_at""", records)
        db.executemany("""INSERT INTO item_assignment_audit (run_id,pair_id,assigned_to,assigned_by,assigned_at)
            VALUES (?,?,?,?,?)""", records)
    refresh_pair_assignments(run_id, pairs)
    return now


def attach_decision_audit(pairs):
    if not pairs:
        return
    ids = [review_pair_id(pair) for pair in pairs]
    history = {key: [] for key in ids}
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        for offset in range(0, len(ids), 900):
            batch = ids[offset:offset + 900]
            slots = ",".join("?" for _ in batch)
            rows = db.execute(
                f"SELECT pair_id, decision, keep_side, canonical_pk, modified_by, modified_at FROM item_decision_audit WHERE pair_id IN ({slots}) ORDER BY id",
                batch,
            )
            for key, decision, keep, canonical, user, modified_at in rows:
                history.setdefault(key, []).append({
                    "decision": decision, "keep": keep,
                    "canonical_pk": json.loads(canonical) if canonical else None,
                    "modified_by": user, "modified_at": modified_at,
                })
    for pair in pairs:
        pair["audit_history"] = history.get(review_pair_id(pair), [])


def create_comparison_run(run_id, tipo, threshold, auto_same_threshold, username):
    now = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        db.execute("INSERT INTO comparison_runs (id,tipo,threshold,auto_same_threshold,state,created_by,created_at) VALUES (?,?,?,?,'queued',?,?)",
                   (run_id, tipo, threshold, auto_same_threshold, username, now))


def update_comparison_run(run_id, state, rows_processed=None, error=None):
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        db.execute("UPDATE comparison_runs SET state=?,rows_processed=COALESCE(?,rows_processed),error=? WHERE id=?",
                   (state, rows_processed, error, run_id))


def persist_comparison_results(run_id, pairs, row_count, candidate_count):
    now = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        db.executemany(
            "INSERT INTO comparison_candidates (run_id,ordinal,pair_id,score,item_a,item_b) VALUES (?,?,?,?,?,?)",
            [(run_id, index, review_pair_id(pair), pair["score"],
              json.dumps(pair["a"], ensure_ascii=False), json.dumps(pair["b"], ensure_ascii=False))
             for index, pair in enumerate(pairs)],
        )
        db.execute("UPDATE comparison_runs SET state='completed',completed_at=?,rows_processed=?,candidate_count=?,retained_count=? WHERE id=?",
                   (now, row_count, candidate_count, len(pairs), run_id))


def comparison_run_list(username, can_manage):
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        db.row_factory = sqlite3.Row
        if not can_manage:
            return [dict(row) for row in db.execute(
                """SELECT r.id,r.tipo,r.threshold,r.state,r.created_by,r.created_at,r.completed_at,
                r.rows_processed,r.candidate_count,(SELECT COUNT(*) FROM item_assignments x
                WHERE x.run_id=r.id AND x.assigned_to=? COLLATE NOCASE) AS retained_count,r.error FROM comparison_runs r
                WHERE r.state!='deleted' AND EXISTS (SELECT 1 FROM item_assignments a
                WHERE a.run_id=r.id AND a.assigned_to=? COLLATE NOCASE) ORDER BY r.created_at DESC LIMIT 100""",
                (username, username),
            )]
        return [dict(row) for row in db.execute(
            "SELECT id,tipo,threshold,state,created_by,created_at,completed_at,rows_processed,candidate_count,retained_count,error FROM comparison_runs WHERE state!='deleted' ORDER BY created_at DESC LIMIT 100"
        )]


def comparison_run_state(run_id):
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        row = db.execute("SELECT state FROM comparison_runs WHERE id=?", (run_id,)).fetchone()
    return row[0] if row else None


def delete_comparison_run(run_id, username):
    now = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        cursor = db.execute("UPDATE comparison_runs SET state='deleted',deleted_by=?,deleted_at=? WHERE id=? AND state='completed'",
                            (username, now, run_id))
        if cursor.rowcount != 1:
            return False
        db.execute("INSERT INTO comparison_run_audit (run_id,action,modified_by,modified_at) VALUES (?,'deleted',?,?)",
                   (run_id, username, now))
    return True


def load_comparison_run(run_id):
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        db.row_factory = sqlite3.Row
        run = db.execute("SELECT * FROM comparison_runs WHERE id=?", (run_id,)).fetchone()
        if not run:
            return None
        result = dict(run)
        if result["state"] == "deleted":
            return None
        if result["state"] == "completed":
            rows = db.execute("SELECT score,item_a,item_b FROM comparison_candidates WHERE run_id=? ORDER BY ordinal", (run_id,))
            result["results"] = [{"score": score, "a": json.loads(item_a), "b": json.loads(item_b)}
                                 for score, item_a, item_b in rows]
            refresh_pair_decisions(result["results"])
        else:
            result["results"] = []
    return result


def load_run_pair(run_id, ordinal):
    with sqlite3.connect(ITEM_REVIEW_DB) as db:
        row = db.execute("SELECT c.score,c.item_a,c.item_b FROM comparison_candidates c JOIN comparison_runs r ON r.id=c.run_id WHERE c.run_id=? AND c.ordinal=? AND r.state='completed'", (run_id, ordinal)).fetchone()
    if not row:
        return None
    pair = {"score": row[0], "a": json.loads(row[1]), "b": json.loads(row[2])}
    refresh_pair_decisions([pair])
    return pair


LEGACY_ASSIGNMENT_HEADERS = ["ORDEN", "ID_CORRIDA"] + [f"{field}_{side}" for side in ("A", "B") for field in (*ITEM_PRIMARY_KEY_FIELDS, "DESCRIPCION", "PRECIO_DE_VENTA")] + ["USUARIO_ORACLE"]
ASSIGNMENT_HEADERS = ["ORDEN", "ID_CORRIDA", "ID_GRUPO"] + LEGACY_ASSIGNMENT_HEADERS[2:]
LEGACY_GROUP_ASSIGNMENT_HEADERS = ["ID_CORRIDA", "ID_GRUPO", "TIPO", "CANTIDAD_ARTICULOS", "SIMILITUD_MINIMA", "SIMILITUD_MAXIMA", "USUARIO_ORACLE"]
PREVIOUS_GROUP_ASSIGNMENT_HEADERS = ["USUARIO_ORACLE", "ID_CORRIDA", "ID_GRUPO", "TIPO", "CANTIDAD_ARTICULOS", "RESUMEN_ARTICULOS", "SIMILITUD_MINIMA", "SIMILITUD_MAXIMA"]
GROUP_ASSIGNMENT_HEADERS = ["USUARIO_ORACLE", "PK_TIPO_REFERENCIA", "PK_SBS_SCC_CODIGO_REFERENCIA", "PK_SBS_CODIGO_REFERENCIA", "PK_CODIGO_REFERENCIA", "ID_GRUPO", "CANTIDAD_ARTICULOS", "RESUMEN_ARTICULOS", "SIMILITUD_MINIMA", "SIMILITUD_MAXIMA", "ID_CORRIDA"]
GROUP_ITEM_HEADERS = ["ID_CORRIDA", "ID_GRUPO", *ITEM_PRIMARY_KEY_FIELDS, "DESCRIPCION", "UNIDAD", "ESTADO", "STOCK1", "PRECIO_DE_VENTA"]


def assignment_workbook(run_id):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName

    run = load_comparison_run(run_id)
    if run is None or run.get("state") != "completed":
        return None
    refresh_pair_assignments(run_id, run["results"])
    groups = build_comparison_groups(run_id, run["results"], run.get("auto_same_threshold", 94))
    try:
        reviewers = oracle_reviewer_users()
    except Exception:
        with sqlite3.connect(ITEM_REVIEW_DB) as db:
            reviewers = [row[0] for row in db.execute(
                "SELECT DISTINCT assigned_to FROM item_group_assignments WHERE run_id=? ORDER BY assigned_to", (run_id,)
            ) if row[0]]
    reviewers = sorted(set(reviewers), key=str.casefold)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Asignaciones"
    details = workbook.create_sheet("Articulos")
    user_sheet = workbook.create_sheet("Usuarios Oracle")
    sheet.append(["ASIGNACIÓN DE GRUPOS · SIS.ITEMS"])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(GROUP_ASSIGNMENT_HEADERS))
    sheet.append(["Una fila representa un grupo completo. Elige un usuario Oracle en USUARIO_ORACLE; los grupos sin usuario conservan su asignación actual."])
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(GROUP_ASSIGNMENT_HEADERS))
    sheet.append(["Solo edita USUARIO_ORACLE (columna A). La clave corresponde al primer artículo de Articulos y ayuda a identificar su bodega o pertenencia; los demás campos son de consulta."])
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(GROUP_ASSIGNMENT_HEADERS))
    sheet.append(GROUP_ASSIGNMENT_HEADERS)
    user_sheet.append(["USUARIO_ORACLE"])
    for reviewer in reviewers:
        user_sheet.append([reviewer])
    if reviewers:
        workbook.defined_names.add(DefinedName("UsuariosOracle", attr_text=f"'Usuarios Oracle'!$A$2:$A${len(reviewers) + 1}"))
        user_sheet.sheet_state = "hidden"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for target, headers in ((details, GROUP_ITEM_HEADERS),):
        cells = []
        for name in headers:
            cell = WriteOnlyCell(target, value=name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cells.append(cell)
        target.append(cells)
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[2].height = 32
    sheet.row_dimensions[3].height = 25
    sheet.row_dimensions[4].height = 48
    sheet[1][0].font = Font(bold=True, size=16, color="FFFFFF")
    sheet[1][0].fill = PatternFill("solid", fgColor="176B55")
    sheet[2][0].font = Font(italic=True, color="365854")
    sheet[2][0].alignment = Alignment(vertical="center", wrap_text=True)
    sheet[3][0].font = Font(italic=True, color="718783")
    sheet[3][0].alignment = Alignment(vertical="center", wrap_text=True)
    for cell in sheet[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for group in groups:
        assigned = [(link.get("assignment") or {}).get("assigned_to") for link in group["links"]]
        assigned_users = {user.casefold(): user for user in assigned if user}
        reviewer = next(iter(assigned_users.values())) if assigned and all(assigned) and len(assigned_users) == 1 else ""
        descriptions = list(dict.fromkeys(str(member["item"].get("DESCRIPCION") or "").strip()
                                          for member in group["members"] if str(member["item"].get("DESCRIPCION") or "").strip()))
        summary = " · ".join(descriptions)
        if len(summary) > 500:
            summary = summary[:497] + "..."
        reference_pk = group["members"][0]["pk"]
        assignment_values = [reviewer]
        assignment_values.extend(reference_pk.get(field) for field in ITEM_PRIMARY_KEY_FIELDS)
        assignment_values.extend((group["id"], len(group["members"]), summary,
                                  group["min_score"], group["max_score"], run_id))
        assignment_cells = []
        for index, value in enumerate(assignment_values):
            cell = WriteOnlyCell(sheet, value=value)
            if index in (0, 1, 2, 3, 4, 5, 10) and value is not None:
                cell.data_type = "s"
                cell.number_format = "@"
            if index in (8, 9):
                cell.number_format = '0.0"%"'
            if index == 7:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            if index == 0:
                cell.protection = Protection(locked=False)
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
                cell.font = Font(bold=True, color="365854")
            assignment_cells.append(cell)
        sheet.append(assignment_cells)
        for member in group["members"]:
            item = member["item"]
            values = [run_id, group["id"]]
            values.extend(item.get(field) for field in ITEM_PRIMARY_KEY_FIELDS)
            values.extend(item.get(field) for field in ("DESCRIPCION", "UNIDAD", "ESTADO", "STOCK1", "PRECIO_DE_VENTA"))
            cells = []
            for value in values:
                cell = WriteOnlyCell(details, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                if GROUP_ITEM_HEADERS[len(cells)] in ("ID_CORRIDA", "ID_GRUPO", *ITEM_PRIMARY_KEY_FIELDS) and value is not None:
                    cell.number_format = "@"
                elif GROUP_ITEM_HEADERS[len(cells)] in ("STOCK1", "PRECIO_DE_VENTA") and isinstance(value, (int, float)):
                    cell.number_format = "#,##0.00"
                cells.append(cell)
            details.append(cells)
    if reviewers and groups:
        validation = DataValidation(type="list", formula1="=UsuariosOracle", allow_blank=True)
        validation.errorTitle = "Usuario Oracle no válido"
        validation.error = "Elige un usuario de la lista desplegable."
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(f"A5:A{len(groups) + 4}")
    sheet.freeze_panes = "B5"
    sheet.auto_filter.ref = f"A4:{excel_column(len(GROUP_ASSIGNMENT_HEADERS))}{len(groups) + 4}"
    for column, width in {"A": 24, "B": 15, "C": 28, "D": 26, "E": 20, "F": 28, "G": 18, "H": 46, "I": 17, "J": 17, "K": 36}.items():
        sheet.column_dimensions[column].width = width
    protection_password = secrets.token_urlsafe(9)
    for protected_sheet in (sheet, details, user_sheet):
        protected_sheet.protection.sheet = True
        protected_sheet.protection.set_password(protection_password)
        protected_sheet.protection.autoFilter = True
        protected_sheet.protection.sort = True
    details.sheet_view.showGridLines = False
    details.freeze_panes = "A2"
    details.auto_filter.ref = f"A1:{excel_column(len(GROUP_ITEM_HEADERS))}{sum(len(group['members']) for group in groups) + 1}"
    for column, width in {"A": 36, "B": 28, "C": 12, "D": 20, "E": 18, "F": 14, "G": 14, "H": 48, "I": 14, "J": 14, "K": 20}.items():
        details.column_dimensions[column].width = width
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def import_assignment_workbook(run_id, contents, assigned_by):
    from openpyxl import load_workbook

    run = load_comparison_run(run_id)
    if run is None or run.get("state") != "completed":
        raise ValueError("La corrida no está disponible para asignar grupos.")
    try:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("El archivo no es un Excel XLSX válido.") from error
    try:
        if "Asignaciones" not in workbook.sheetnames:
            raise ValueError("Falta la hoja Asignaciones.")
        sheet = workbook["Asignaciones"]
        header = None
        header_row = None
        for line_number, candidate in enumerate(sheet.iter_rows(values_only=True), 1):
            candidate = list(candidate or [])
            if candidate in (ASSIGNMENT_HEADERS, LEGACY_ASSIGNMENT_HEADERS):
                header = candidate
                workbook.close()
                return import_legacy_assignment_workbook(run_id, contents, assigned_by)
            if candidate in (GROUP_ASSIGNMENT_HEADERS, PREVIOUS_GROUP_ASSIGNMENT_HEADERS, LEGACY_GROUP_ASSIGNMENT_HEADERS):
                header, header_row = candidate, line_number
                break
        if header is None:
            raise ValueError("No se encontró el encabezado de asignación por grupos.")
        if "Articulos" not in workbook.sheetnames or list(next(workbook["Articulos"].iter_rows(values_only=True), None) or []) != GROUP_ITEM_HEADERS:
            raise ValueError("Falta la hoja Articulos o sus columnas no coinciden con la plantilla.")
        try:
            known_users = {user.casefold() for user in oracle_reviewer_users()}
        except Exception as error:
            raise ValueError("No se pudo verificar el catálogo de usuarios Oracle.") from error
        positions = {name: index for index, name in enumerate(header)}
        groups = build_comparison_groups(run_id, run["results"], run.get("auto_same_threshold", 94))
        group_by_id = {group["id"]: group for group in groups}
        assignments, seen, errors, unchanged = {}, set(), [], 0
        rows = sheet.iter_rows(min_row=header_row + 1, values_only=True)
        for row_number, row in enumerate(rows, header_row + 1):
            if not row or all(value is None or value == "" for value in row):
                continue
            supplied_run = str(row[positions["ID_CORRIDA"]] or "").strip() if len(row) > positions["ID_CORRIDA"] else ""
            group_id = str(row[positions["ID_GRUPO"]] or "").strip() if len(row) > positions["ID_GRUPO"] else ""
            username = row[positions["USUARIO_ORACLE"]] if len(row) > positions["USUARIO_ORACLE"] else None
            if supplied_run != run_id:
                errors.append(f"fila {row_number}: ID_CORRIDA no corresponde a esta corrida")
                continue
            if group_id not in group_by_id:
                errors.append(f"fila {row_number}: ID_GRUPO no pertenece a esta corrida")
                continue
            if group_id in seen:
                errors.append(f"fila {row_number}: ID_GRUPO duplicado; debe existir una sola fila por grupo")
                continue
            seen.add(group_id)
            if username is None or str(username).strip() == "":
                unchanged += 1
                continue
            username = str(username).strip()
            if not IDENTIFIER.fullmatch(username):
                errors.append(f"fila {row_number}: usuario Oracle inválido")
                continue
            if username.casefold() not in known_users:
                errors.append(f"fila {row_number}: el usuario no existe en Oracle")
                continue
            assignments[group_id] = username
        if errors:
            raise ValueError("No se guardó ninguna asignación. " + "; ".join(errors[:10]))
        if not assignments:
            raise ValueError("Elige un usuario Oracle en USUARIO_ORACLE para al menos un grupo.")
        pair_by_id = {review_pair_id(pair): pair for pair in run["results"]}
        pair_assignments = {}
        for group_id, reviewer in assignments.items():
            for link in group_by_id[group_id]["links"]:
                pair_assignments[link["pair_id"]] = (pair_by_id[link["pair_id"]], reviewer)
        now = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
        records = [(run_id, pair_id, reviewer, assigned_by, now)
                   for pair_id, (_pair, reviewer) in pair_assignments.items()]
        group_records = [(run_id, group_id, reviewer, assigned_by, now)
                         for group_id, reviewer in assignments.items()]
        with sqlite3.connect(ITEM_REVIEW_DB) as db:
            db.execute("BEGIN IMMEDIATE")
            db.executemany("""INSERT INTO item_assignments (run_id,pair_id,assigned_to,assigned_by,assigned_at)
                VALUES (?,?,?,?,?) ON CONFLICT(run_id,pair_id) DO UPDATE SET
                assigned_to=excluded.assigned_to,assigned_by=excluded.assigned_by,assigned_at=excluded.assigned_at""", records)
            db.executemany("INSERT INTO item_assignment_audit (run_id,pair_id,assigned_to,assigned_by,assigned_at) VALUES (?,?,?,?,?)", records)
            db.executemany("""INSERT INTO item_group_assignments (run_id,group_id,assigned_to,assigned_by,assigned_at)
                VALUES (?,?,?,?,?) ON CONFLICT(run_id,group_id) DO UPDATE SET
                assigned_to=excluded.assigned_to,assigned_by=excluded.assigned_by,assigned_at=excluded.assigned_at""", group_records)
            db.executemany("INSERT INTO item_group_assignment_audit (run_id,group_id,assigned_to,assigned_by,assigned_at) VALUES (?,?,?,?,?)", group_records)
        refresh_pair_assignments(run_id, [pair for pair, _reviewer in pair_assignments.values()])
        return {"assigned_count": len(assignments), "assigned_pairs_count": len(records), "unchanged_count": unchanged}
    finally:
        try: workbook.close()
        except Exception: pass


def import_legacy_assignment_workbook(run_id, contents, assigned_by):
    from openpyxl import load_workbook

    run = load_comparison_run(run_id)
    if run is None or run.get("state") != "completed":
        raise ValueError("La corrida no está disponible para asignar pares.")
    try:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("El archivo no es un Excel XLSX válido.") from error
    try:
        if "Asignaciones" not in workbook.sheetnames:
            raise ValueError("Falta la hoja Asignaciones.")
        sheet = workbook["Asignaciones"]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        header = list(header or [])
        if header not in (ASSIGNMENT_HEADERS, LEGACY_ASSIGNMENT_HEADERS):
            raise ValueError("La plantilla no coincide con el formato esperado.")
        positions = {name: index for index, name in enumerate(header)}
        group_by_pair, group_by_id = {}, {}
        pair_by_id = {review_pair_id(pair): pair for pair in run["results"]}
        if "ID_GRUPO" in positions:
            for group in build_comparison_groups(run_id, run["results"], run.get("auto_same_threshold", 94)):
                group_by_id[group["id"]] = group
                for link in group["links"]:
                    group_by_pair[link["pair_id"]] = group["id"]
        selected, seen, errors, unchanged = [], set(), [], 0
        group_assignments, blank_group_rows = {}, {}
        for row_number, row in enumerate(rows, 2):
            if not row or all(value is None or value == "" for value in row):
                continue
            username = row[positions["USUARIO_ORACLE"]] if len(row) > positions["USUARIO_ORACLE"] else None
            if username is None or str(username).strip() == "":
                unchanged += 1
                if "ID_GRUPO" in positions:
                    blank_group = str(row[positions["ID_GRUPO"]] or "").strip()
                    if blank_group:
                        blank_group_rows[blank_group] = blank_group_rows.get(blank_group, 0) + 1
                continue
            username = str(username).strip()
            ordinal_raw = row[positions["ORDEN"]] if len(row) > positions["ORDEN"] else None
            if not IDENTIFIER.fullmatch(username):
                errors.append(f"fila {row_number}: usuario Oracle inválido")
                continue
            if str(row[positions["ID_CORRIDA"]]).strip() != run_id:
                errors.append(f"fila {row_number}: ID_CORRIDA no corresponde a esta corrida")
                continue
            try:
                ordinal = int(ordinal_raw)
                if isinstance(ordinal_raw, bool) or float(ordinal_raw) != ordinal:
                    raise ValueError
            except (ValueError, TypeError):
                errors.append(f"fila {row_number}: ORDEN inválido")
                continue
            if not 1 <= ordinal <= len(run["results"]) or ordinal in seen:
                errors.append(f"fila {row_number}: ORDEN duplicado o fuera de rango")
                continue
            pair = run["results"][ordinal - 1]
            valid = True
            for side in ("A", "B"):
                candidate = pair[side.lower()]
                for field in ITEM_PRIMARY_KEY_FIELDS:
                    index = positions[f"{field}_{side}"]
                    actual = row[index] if len(row) > index else None
                    if str(actual if actual is not None else "").strip() != str(candidate.get(field) if candidate.get(field) is not None else "").strip():
                        valid = False
            if not valid:
                errors.append(f"fila {row_number}: las claves A/B no coinciden con el par guardado")
                continue
            supplied_group = ""
            if "ID_GRUPO" in positions:
                supplied_group = str(row[positions["ID_GRUPO"]] or "").strip()
                expected_group = group_by_pair.get(review_pair_id(pair), "")
                if str(supplied_group or "").strip() != expected_group:
                    errors.append(f"fila {row_number}: ID_GRUPO no corresponde a las claves A/B")
                    continue
                if supplied_group:
                    previous_user = group_assignments.get(supplied_group)
                    if previous_user and previous_user.casefold() != username.casefold():
                        errors.append(f"fila {row_number}: todas las relaciones de un grupo deben tener el mismo usuario Oracle")
                        continue
                    group_assignments[supplied_group] = username
            seen.add(ordinal)
            pair["_assigned_to"] = username
            selected.append(pair)
        if errors:
            raise ValueError("No se guardó ninguna asignación. " + "; ".join(errors[:10]))
        if group_assignments:
            selected_by_pair = {review_pair_id(pair): pair for pair in selected}
            for group_id, reviewer in group_assignments.items():
                group = group_by_id[group_id]
                for link in group["links"]:
                    pair = pair_by_id.get(link["pair_id"])
                    if pair is not None:
                        pair["_assigned_to"] = reviewer
                        selected_by_pair[link["pair_id"]] = pair
                unchanged -= blank_group_rows.get(group_id, 0)
            selected = list(selected_by_pair.values())
        if not selected:
            raise ValueError("El Excel no contiene usuarios en USUARIO_ORACLE.")
        # Each row may target a different reviewer. Persist all rows and audit in one transaction.
        now = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
        records = [(run_id, review_pair_id(pair), pair["_assigned_to"], assigned_by, now) for pair in selected]
        with sqlite3.connect(ITEM_REVIEW_DB) as db:
            db.execute("BEGIN IMMEDIATE")
            db.executemany("""INSERT INTO item_assignments (run_id,pair_id,assigned_to,assigned_by,assigned_at)
                VALUES (?,?,?,?,?) ON CONFLICT(run_id,pair_id) DO UPDATE SET
                assigned_to=excluded.assigned_to,assigned_by=excluded.assigned_by,assigned_at=excluded.assigned_at""", records)
            db.executemany("INSERT INTO item_assignment_audit (run_id,pair_id,assigned_to,assigned_by,assigned_at) VALUES (?,?,?,?,?)", records)
        return {"assigned_count": len(selected), "unchanged_count": unchanged}
    finally:
        workbook.close()


def save_item_decision(pair, decision, keep, run_id):
    if decision not in ("pending", "same", "different") or (decision == "same" and keep not in ("a", "b")) or (decision != "same" and keep not in (None, "")):
        raise ValueError("Selecciona pendiente, mismo artículo o artículos diferentes; si es el mismo, elige la clave que se conservará.")
    canonical = {field: pair[keep][field] for field in ITEM_PRIMARY_KEY_FIELDS} if decision == "same" else None
    record = {
        "decision": decision,
        "keep": keep if decision == "same" else None,
        "canonical_pk": canonical,
        "price_difference": calculate_price_difference(pair),
        "saved_by": getattr(request_context, "username", "desconocido"),
        "saved_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
    }
    with items_review_lock:
        with sqlite3.connect(ITEM_REVIEW_DB) as db:
            db.execute("""INSERT INTO item_decisions
                (pair_id, decision, keep_side, canonical_pk, saved_by, saved_at)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(pair_id) DO UPDATE SET decision=excluded.decision,
                keep_side=excluded.keep_side, canonical_pk=excluded.canonical_pk,
                saved_by=excluded.saved_by, saved_at=excluded.saved_at""",
                (review_pair_id(pair), decision, record["keep"],
                 json.dumps(canonical, ensure_ascii=False) if canonical else None,
                 record["saved_by"], record["saved_at"]))
            db.execute("""INSERT INTO item_decision_audit
                (run_id,pair_id,decision,keep_side,canonical_pk,modified_by,modified_at)
                VALUES (?,?,?,?,?,?,?)""",
                (run_id, review_pair_id(pair), decision, record["keep"],
                 json.dumps(canonical, ensure_ascii=False) if canonical else None,
                 record["saved_by"], record["saved_at"]))
    return record


def run_items_export(job_id, tipo, expected_rows, output_dir, credentials):
    global items_export_running
    started = time.monotonic()
    parquet_dir = output_dir / "parquet"
    parquet_dir.mkdir(parents=True, exist_ok=True)
    env = {
        **os.environ,
        "ORACLE_SCHEMA": "SIS",
        "ORACLE_TABLE": "ITEMS",
        "ORACLE_FILTER_COLUMN": "TIPO",
        "ORACLE_FILTER_VALUE": tipo,
        "ORACLE_MAX_ROWS": "all",
        "ORACLE_FETCH_ARRAY_SIZE": "1000",
        "EXPORT_PROGRESS": "1",
        "PARQUET_OUTPUT_DIR": str(parquet_dir),
        "ORACLE_USER": credentials[0], "ORACLE_PASSWORD": credentials[1],
    }
    try:
        with items_jobs_lock:
            items_jobs[job_id]["state"] = "reading"
        process = subprocess.Popen(
            ["bash", str(PROBE)], cwd=ROOT, env=env,
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
            text=True, bufsize=1,
        )
        output_tail = ""
        for line in process.stdout:
            output_tail = (output_tail + line)[-MAX_OUTPUT:]
            progress = re.search(r"EXPORT_PROGRESS rows=(\d+) batch=(\d+) state=([a-z_]+)", line)
            if progress:
                with items_jobs_lock:
                    items_jobs[job_id].update({
                        "rows_processed": int(progress.group(1)),
                        "current_batch": int(progress.group(2)),
                        "state": progress.group(3),
                    })
        return_code = process.wait()
        exported = EXPORT_RE.search(output_tail)
        verified = VERIFY_RE.search(output_tail)
        parquet_file = parquet_dir / "part-00001.parquet"
        if return_code != 0 or not exported or not verified or not parquet_file.is_file():
            raise RuntimeError("Falló la extracción Zig o la verificación Parquet.")
        rows = int(exported.group(3))
        excel_path = output_dir / "SIS_ITEMS_TIPO.xlsx"
        with items_jobs_lock:
            items_jobs[job_id].update({"state": "writing", "rows_processed": rows})
        convert_items_parquet_to_excel(parquet_file, excel_path, tipo, rows)
        shutil.rmtree(parquet_dir, ignore_errors=True)
        with items_jobs_lock:
            items_jobs[job_id].update({
                "state": "completed",
                "rows_processed": rows,
                "columns": int(verified.group(2)),
                "bytes_written": excel_path.stat().st_size,
                "elapsed_ms": round((time.monotonic() - started) * 1000),
                "output_file": str(excel_path.relative_to(ROOT)),
                "download_url": f"/api/v1/items/exports/{job_id}/download",
            })
    except Exception as exc:
        shutil.rmtree(output_dir, ignore_errors=True)
        message = str(exc) if isinstance(exc, ValueError) else "No se pudo completar la exportación del tipo seleccionado."
        with items_jobs_lock:
            if job_id in items_jobs:
                items_jobs[job_id].update({"state": "failed", "error": message})
    finally:
        with items_jobs_lock:
            items_export_running = False


def run_items_compare(job_id, tipo, threshold, expected_rows, output_dir, credentials):
    global items_compare_running
    started = time.monotonic()
    parquet_dir = output_dir / "parquet"
    parquet_dir.mkdir(parents=True, exist_ok=True)
    env = {
        **os.environ,
        "ORACLE_SCHEMA": "SIS",
        "ORACLE_TABLE": "ITEMS",
        "ORACLE_FILTER_COLUMN": "TIPO",
        "ORACLE_FILTER_VALUE": tipo,
        "ORACLE_MAX_ROWS": "all",
        "ORACLE_FETCH_ARRAY_SIZE": "1000",
        "EXPORT_PROGRESS": "1",
        "PARQUET_OUTPUT_DIR": str(parquet_dir),
        "ORACLE_USER": credentials[0], "ORACLE_PASSWORD": credentials[1],
    }
    try:
        with items_jobs_lock:
            items_compare_jobs[job_id]["state"] = "reading"
        update_comparison_run(job_id, "reading")
        process = subprocess.Popen(
            ["bash", str(PROBE)], cwd=ROOT, env=env,
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
            text=True, bufsize=1,
        )
        output_tail = ""
        for line in process.stdout:
            output_tail = (output_tail + line)[-MAX_OUTPUT:]
            progress = re.search(r"EXPORT_PROGRESS rows=(\d+) batch=(\d+) state=([a-z_]+)", line)
            if progress:
                with items_jobs_lock:
                    export_state = progress.group(3)
                    if export_state == "completed":
                        export_state = "reading"
                    items_compare_jobs[job_id].update({
                        "rows_processed": int(progress.group(1)),
                        "current_batch": int(progress.group(2)),
                        "state": export_state,
                    })
        return_code = process.wait()
        exported = EXPORT_RE.search(output_tail)
        verified = VERIFY_RE.search(output_tail)
        parquet_file = parquet_dir / "part-00001.parquet"
        if return_code != 0 or not exported or not verified or not parquet_file.is_file():
            raise RuntimeError("Falló la extracción Zig o la verificación Parquet.")
        with items_jobs_lock:
            items_compare_jobs[job_id].update({"state": "comparing", "rows_processed": int(exported.group(3))})
        update_comparison_run(job_id, "comparing", int(exported.group(3)))
        import pyarrow.parquet as pq
        columns = list(dict.fromkeys((*ITEM_PRIMARY_KEY_FIELDS, "DESCRIPCION", "UNIDAD", "ESTADO", "STOCK1", "STOCK2", "STOCK_COMPROMETIDO", "COSTO_PROMEDIO", "PRECIO_DE_VENTA")))
        records = pq.read_table(parquet_file, columns=columns).to_pylist()
        pairs, candidate_count = compare_item_records(records, threshold)
        persist_comparison_results(job_id, pairs, len(records), candidate_count)
        pairs = attach_item_decisions(pairs)
        attach_decision_audit(pairs)
        shutil.rmtree(output_dir, ignore_errors=True)
        with items_jobs_lock:
            items_compare_jobs[job_id].update({
                "state": "completed",
                "rows_processed": len(records),
                "candidate_count": candidate_count,
                "retained_count": len(pairs),
                "results": pairs,
                "elapsed_ms": round((time.monotonic() - started) * 1000),
            })
    except Exception as exc:
        shutil.rmtree(output_dir, ignore_errors=True)
        message = str(exc) if isinstance(exc, ValueError) else "No se pudo completar la comparación de descripciones."
        with items_jobs_lock:
            if job_id in items_compare_jobs:
                items_compare_jobs[job_id].update({"state": "failed", "error": message})
        update_comparison_run(job_id, "failed", error=message)
    finally:
        with items_jobs_lock:
            items_compare_running = False


def cleanup_expired_downloads():
    now = time.time()
    with downloads_lock:
        expired = [token for token, record in downloads.items() if record["expires_at"] <= now]
        for token in expired:
            shutil.rmtree(downloads.pop(token)["directory"], ignore_errors=True)
        active_directories = {record["directory"] for record in downloads.values()}
    temp_root = Path(tempfile.gettempdir())
    for directory in temp_root.glob("zig-proyecto-parquet-download-*"):
        try:
            if directory not in active_directories and now - directory.stat().st_mtime >= DOWNLOAD_TTL:
                shutil.rmtree(directory, ignore_errors=True)
        except FileNotFoundError:
            pass


def validate_oracle_login(username, password):
    host, sid = os.environ.get("ORACLE_HOST", ""), os.environ.get("ORACLE_SID", "")
    if not host or not sid or not username or not password:
        return False
    dsn = oracledb.makedsn(host, int(os.environ.get("ORACLE_PORT", "1521")), sid=sid)
    try:
        with oracledb.connect(user=username, password=password, dsn=dsn):
            return True
    except Exception:
        return False


def session_for(handler):
    cookie = handler.headers.get("Cookie", "")
    match = re.search(r"(?:^|;\s*)zig_session=([A-Za-z0-9_-]{40,60})(?:;|$)", cookie)
    if not match:
        return None, None
    token = match.group(1)
    with sessions_lock:
        session = sessions.get(token)
        if session and session["expires_at"] > time.time():
            session["expires_at"] = time.time() + SESSION_TTL
            return token, session
        sessions.pop(token, None)
    return None, None


class Handler(BaseHTTPRequestHandler):
    server_version = "OracleParquetProbe/1.0"

    def log_message(self, _format, *_args):
        # Evita guardar rutas/errores de los trabajos en logs innecesarios.
        return

    def send_json(self, status, value):
        payload = json.dumps(value, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(payload)

    def valid_local_request(self):
        host = self.headers.get("Host", "")
        if host not in (f"{HOST}:{PORT}", f"localhost:{PORT}"):
            return False
        origin = self.headers.get("Origin")
        return not origin or origin in (f"http://{HOST}:{PORT}", f"http://localhost:{PORT}")

    def do_GET(self):
        cleanup_expired_downloads()
        if not self.valid_local_request():
            self.send_error(403)
            return
        path = unquote(urlsplit(self.path).path)
        if path in ("/login", "/login.html") or path.startswith("/css/"):
            if path == "/login":
                path = "/login.html"
        else:
            _token, session = session_for(self)
            if not session:
                if path.startswith("/api/"):
                    self.send_json(401, {"error": "Inicia sesión con tu usuario Oracle."})
                else:
                    self.send_response(302)
                    self.send_header("Location", "/login")
                    self.end_headers()
                return
            request_context.credentials = session["credentials"]
            request_context.username = session["username"]
        if path.startswith("/api/"):
            comparison_get = (
                path in ("/api/v1/session", "/api/v1/items/types", "/api/v1/items/compare/runs")
                or re.fullmatch(r"/api/v1/items/compare/[0-9a-f]{32}(?:/groups|/assignments/template|/merge-plan)?", path)
            )
            if not comparison_get:
                self.send_json(404, {"error": "Esta aplicación está habilitada solo para comparar descripciones."})
                return
        if path == "/api/v1/items/compare/runs":
            try:
                username = request_context.username
                can_manage = username.casefold() == COMPARISON_RUNNER_USER
                self.send_json(200, {"runs": comparison_run_list(username, can_manage)})
            except sqlite3.Error:
                self.send_json(503, {"error": "No se pudo leer el historial de corridas."})
            return
        if path == "/api/v1/session":
            can_manage = request_context.username.casefold() == COMPARISON_RUNNER_USER
            self.send_json(200, {
                "username": request_context.username,
                "can_run_comparisons": can_manage,
                "can_delete_comparisons": can_manage,
            })
            return
        groups_match = re.fullmatch(r"/api/v1/items/compare/([0-9a-f]{32})/groups", path)
        if groups_match:
            query = parse_qs(urlsplit(self.path).query)
            try:
                page = int(query.get("page", ["1"])[0])
            except ValueError:
                page = 0
            decision_filter = query.get("decision", ["pending"])[0]
            if decision_filter not in ("pending", "unify", "skip", "manual", "all"):
                decision_filter = "pending"
            self.items_compare_groups(groups_match.group(1), page, decision_filter)
            return
        merge_plan_match = re.fullmatch(r"/api/v1/items/compare/([0-9a-f]{32})/merge-plan", path)
        if merge_plan_match:
            if request_context.username.casefold() != COMPARISON_RUNNER_USER:
                self.send_json(403, {"error": "Solo la cuenta responsable puede descargar el plan de empates."})
                return
            contents = group_merge_plan_workbook(merge_plan_match.group(1))
            if contents is None:
                self.send_json(404, {"error": "La corrida no está disponible."})
                return
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f"attachment; filename=SIS_ITEMS_PLAN_EMPATES_{merge_plan_match.group(1)}.xlsx")
            self.send_header("Content-Length", str(len(contents)))
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.end_headers()
            self.wfile.write(contents)
            return
        assignment_template_match = re.fullmatch(r"/api/v1/items/compare/([0-9a-f]{32})/assignments/template", path)
        if assignment_template_match:
            if request_context.username.casefold() != COMPARISON_RUNNER_USER:
                self.send_json(403, {"error": "Solo la cuenta responsable puede descargar la plantilla."})
                return
            contents = assignment_workbook(assignment_template_match.group(1))
            if contents is None:
                self.send_json(404, {"error": "La corrida no está disponible."})
                return
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f"attachment; filename=SIS_ITEMS_ASIGNACIONES_{assignment_template_match.group(1)}.xlsx")
            self.send_header("Content-Length", str(len(contents)))
            self.send_header("X-Content-Type-Options", "nosniff")
            self.end_headers()
            self.wfile.write(contents)
            return
        if path == "/api/v1/files/SIS_ITEMS_TIPO_I.xlsx":
            self.download_sis_items_excel()
            return
        if path == "/api/v1/items/types":
            self.catalog_item_types()
            return
        match = re.fullmatch(r"/api/v1/items/compare/([0-9a-f]{32})", path)
        if match:
            query = parse_qs(urlsplit(self.path).query)
            page_raw = query.get("page", ["1"])[0]
            try:
                page = int(page_raw)
            except ValueError:
                page = 0
            filters = {field: query.get(f"f_{field}", [""])[0][:80].strip()
                       for field in ITEM_PRIMARY_KEY_FIELDS}
            decision_filter = query.get("decision", [""])[0]
            filters["decision"] = decision_filter if decision_filter in ("all", "pending", "same", "different") else "pending"
            self.items_compare_status(match.group(1), page, filters)
            return
        match = re.fullmatch(r"/api/v1/items/exports/([0-9a-f]{32})(/download)?", path)
        if match:
            if match.group(2):
                self.download_items_excel(match.group(1))
            else:
                self.items_export_status(match.group(1))
            return
        match = re.fullmatch(r"/api/v1/downloads/([A-Za-z0-9_-]{40,50})", path)
        if match:
            self.download_parquet(match.group(1))
            return
        match = re.fullmatch(r"/api/v1/full-exports/([0-9a-f]{32})(/download)?", path)
        if match:
            if match.group(2):
                self.download_full_export(match.group(1))
            else:
                self.full_export_status(match.group(1))
            return
        if path == "/api/v1/schemas":
            self.catalog_schemas()
            return
        match = re.fullmatch(r"/api/v1/schemas/([^/]+)/tables(?:/([^/]+))?", path)
        if match:
            schema, table = match.groups()
            if not valid_identifier(schema) or (table is not None and not valid_identifier(table)):
                self.send_json(400, {"error": "Identificador Oracle no válido."})
            elif table is None:
                self.catalog_tables(schema)
            else:
                self.catalog_columns(schema, table)
            return
        if path in ("/", "/items/compare", "/items/review", "/items-compare.html"):
            path = "/items-compare.html"
        elif path in ("/probe", "/probe.html", "/items", "/items.html", "/index.html"):
            self.send_response(302)
            self.send_header("Location", "/items/compare")
            self.end_headers()
            return
        elif path.endswith(".html") and path != "/login.html":
            self.send_error(404)
            return
        candidate = (WEB / path.lstrip("/")).resolve()
        if WEB.resolve() not in candidate.parents or not candidate.is_file():
            self.send_error(404)
            return
        guessed_type = mimetypes.guess_type(candidate.name)[0] or "application/octet-stream"
        content_type = f"{guessed_type}; charset=utf-8" if guessed_type.startswith(("text/", "application/javascript")) else guessed_type
        payload = candidate.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(payload)))
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(payload)

    def download_sis_items_excel(self):
        file_path = SIS_ITEMS_EXCEL.resolve()
        if ROOT not in file_path.parents or not file_path.is_file():
            self.send_json(404, {"error": "El Excel SIS.ITEMS tipo I no está disponible."})
            return
        try:
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Length", str(file_path.stat().st_size))
            self.send_header("Content-Disposition", 'attachment; filename="SIS_ITEMS_TIPO_I.xlsx"')
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.end_headers()
            with file_path.open("rb") as excel_file:
                shutil.copyfileobj(excel_file, self.wfile, length=1024 * 1024)
        except (BrokenPipeError, ConnectionResetError):
            pass

    def catalog_item_types(self):
        try:
            self.send_json(200, {"items": item_type_counts()})
        except Exception:
            self.send_json(503, {"error": "No se pudo cargar el catálogo de tipos de SIS.ITEMS."})

    def items_export_status(self, job_id):
        with items_jobs_lock:
            job = items_jobs.get(job_id)
            if job is None:
                self.send_json(404, {"error": "No se encontró la exportación."})
                return
            response = {key: value for key, value in job.items() if key != "output_file"}
        self.send_json(200, response)

    def items_compare_status(self, job_id, page=1, filters=None):
        if comparison_run_state(job_id) in (None, "deleted"):
            self.send_json(404, {"error": "La corrida no existe o fue retirada del historial."})
            return
        with items_jobs_lock:
            job = items_compare_jobs.get(job_id)
        if job is None:
            job = load_comparison_run(job_id)
        if job is None:
            self.send_json(404, {"error": "No se encontró esta comparación."})
            return
        with items_jobs_lock:
            if job.get("state") == "completed":
                refresh_pair_decisions(job["results"])
                refresh_pair_assignments(job_id, job["results"])
            response = {key: value for key, value in job.items() if key != "results"}
            if job.get("state") == "completed":
                page_size = 200
                filters = filters or {}
                filtered = []
                can_manage = request_context.username.casefold() == COMPARISON_RUNNER_USER
                visible_pairs = [pair for pair in job["results"] if can_manage or
                                 (pair.get("assignment") or {}).get("assigned_to", "").casefold() == request_context.username.casefold()]
                resolved_count = sum(1 for pair in visible_pairs if pair.get("decision") in ("same", "different"))
                for index, pair in enumerate(job["results"]):
                    if not can_manage and (pair.get("assignment") or {}).get("assigned_to", "").casefold() != request_context.username.casefold():
                        continue
                    if filters.get("decision") not in (None, "", "all") and (pair.get("decision") or "pending") != filters["decision"]:
                        continue
                    matches = True
                    for field in ITEM_PRIMARY_KEY_FIELDS:
                        needle = filters.get(field, "").casefold()
                        if needle and not any(needle in str(pair[side].get(field) or "").casefold() for side in ("a", "b")):
                            matches = False
                            break
                    if matches:
                        filtered.append((index, pair))
                filtered_count = len(filtered)
                page_count = max(1, (filtered_count + page_size - 1) // page_size)
                if page < 1 or page > page_count:
                    self.send_json(400, {"error": "La página solicitada no existe."})
                    return
                start = (page - 1) * page_size
                attach_decision_audit([pair for _, pair in filtered[start:start + page_size]])
                response.update({
                    "results": [{**pair, "review_index": index} for index, pair in filtered[start:start + page_size]],
                    "page": page,
                    "page_size": page_size,
                    "page_count": page_count,
                    "shown_count": min(page_size, max(0, filtered_count - start)),
                    "filtered_count": filtered_count,
                    "review_total": len(visible_pairs),
                    "review_resolved": resolved_count,
                    "review_pending": len(visible_pairs) - resolved_count,
                    "results_truncated": job["candidate_count"] > job["retained_count"],
                })
        self.send_json(200, response)

    def items_compare_groups(self, run_id, page=1, decision_filter="pending"):
        if comparison_run_state(run_id) in (None, "deleted"):
            self.send_json(404, {"error": "La corrida no existe o fue retirada del historial."})
            return
        run = load_comparison_run(run_id)
        if run is None or run.get("state") != "completed":
            self.send_json(404, {"error": "La corrida no está lista para revisar grupos."})
            return
        refresh_pair_assignments(run_id, run["results"])
        groups = build_comparison_groups(run_id, run["results"], run.get("auto_same_threshold", 94))
        attach_group_decisions(run_id, groups)
        username = request_context.username
        can_manage = username.casefold() == COMPARISON_RUNNER_USER
        visible = []
        for group in groups:
            best_link_score = {}
            for link in group["links"]:
                for item_key in (item_pk_tuple(link["a_pk"]), item_pk_tuple(link["b_pk"])):
                    best_link_score[item_key] = max(best_link_score.get(item_key, 0), link["score"])
            canonical_tuple = None
            group["members"].sort(key=lambda member: (
                0 if canonical_tuple and item_pk_tuple(member["pk"]) == canonical_tuple else 1,
                -best_link_score.get(item_pk_tuple(member["pk"]), 0),
                item_pk_tuple(member["pk"]),
            ))
            assignees = [(link.get("assignment") or {}).get("assigned_to") for link in group["links"]]
            consistent_assignment = bool(assignees) and all(
                assignee and assignee.casefold() == assignees[0].casefold() for assignee in assignees
            )
            group["assigned_to"] = assignees[0] if consistent_assignment else None
            group["assignment_complete"] = consistent_assignment
            group["assigned_by"] = (group["links"][0].get("assignment") or {}).get("assigned_by") if consistent_assignment else None
            group["assigned_at"] = (group["links"][0].get("assignment") or {}).get("assigned_at") if consistent_assignment else None
            group["can_review"] = consistent_assignment and assignees[0].casefold() == username.casefold()
            if can_manage or (consistent_assignment and assignees[0].casefold() == username.casefold()):
                visible.append(group)
        has_confirmed_selection = lambda group: bool(group.get("subgroups"))
        has_pending_members = lambda group: bool(group.get("pending_keys"))
        has_skipped_members = lambda group: bool(group.get("skipped_members"))
        counts = {
            "pending": sum(1 for group in visible if has_pending_members(group)),
            "unify": sum(1 for group in visible if has_confirmed_selection(group)),
            "skip": sum(1 for group in visible if has_skipped_members(group)),
            "manual": 0,
        }
        if decision_filter == "pending":
            filtered = [group for group in visible if has_pending_members(group)]
        elif decision_filter == "unify":
            # A component can contain a confirmed same-item subset and still
            # have other members awaiting review. Keep it visible in both filters.
            filtered = [group for group in visible if has_confirmed_selection(group)]
        elif decision_filter == "skip":
            filtered = [group for group in visible if has_skipped_members(group)]
        else:
            filtered = visible
        page_size = 25
        page_count = max(1, (len(filtered) + page_size - 1) // page_size)
        if page < 1 or page > page_count:
            self.send_json(400, {"error": "La página solicitada no existe."})
            return
        start = (page - 1) * page_size
        page_groups = []
        for group in filtered[start:start + page_size]:
            group.pop("member_keys", None)
            page_groups.append(group)
        self.send_json(200, {
            "id": run_id, "tipo": run["tipo"], "auto_same_threshold": run.get("auto_same_threshold", 94),
            "groups": page_groups, "page": page, "page_size": page_size, "page_count": page_count,
            "filtered_count": len(filtered), "group_total": len(visible), "group_pending": counts["pending"],
            "group_unify": counts["unify"], "group_manual": counts["manual"],
            "group_skip": counts["skip"],
            "subgroup_total": sum(len(group.get("subgroups", [])) for group in visible),
            "hidden_assignment_groups": len(groups) - len(visible),
        })

    def save_items_compare_group_decision(self, run_id, payload):
        if comparison_run_state(run_id) != "completed":
            self.send_json(404, {"error": "La corrida no está disponible para guardar grupos."})
            return
        group_id, decision = payload.get("group_id"), payload.get("decision")
        canonical_pk = payload.get("canonical_pk")
        member_pks = payload.get("member_pks")
        if not isinstance(group_id, str) or not re.fullmatch(r"[0-9a-f]{24}", group_id) or decision not in ("unify", "skip", "pending"):
            self.send_json(400, {"error": "El grupo o la decisión no son válidos."})
            return
        run = load_comparison_run(run_id)
        if run is None:
            self.send_json(404, {"error": "No se encontró la corrida."})
            return
        refresh_pair_assignments(run_id, run["results"])
        groups = build_comparison_groups(run_id, run["results"], run.get("auto_same_threshold", 94))
        group = next((candidate for candidate in groups if candidate["id"] == group_id), None)
        if group is None:
            self.send_json(404, {"error": "El grupo ya no coincide con los candidatos de la corrida."})
            return
        username = request_context.username
        assignees = [(link.get("assignment") or {}).get("assigned_to") for link in group["links"]]
        if not assignees or any(not assignee or assignee.casefold() != username.casefold() for assignee in assignees):
            self.send_json(403, {"error": "El grupo completo debe estar asignado a tu usuario Oracle."})
            return
        minimum_members = 2 if decision == "unify" else 1
        if not isinstance(member_pks, list) or len(member_pks) < minimum_members:
            self.send_json(400, {"error": "Selecciona al menos dos artículos para confirmar que son el mismo, o uno para marcarlo como omitido."})
            return
        member_by_key = {item_pk_tuple(member["pk"]): member["pk"] for member in group["members"]}
        try:
            selected_keys = [item_pk_tuple(pk) for pk in member_pks]
        except Exception:
            selected_keys = []
        if len(set(selected_keys)) != len(selected_keys) or any(key not in member_by_key for key in selected_keys):
            self.send_json(400, {"error": "La selección contiene claves repetidas o ajenas al grupo."})
            return
        selected_keys = sorted(selected_keys)
        member_pks = [member_by_key[key] for key in selected_keys]
        members_json = json.dumps(member_pks, ensure_ascii=False, sort_keys=True)
        subgroup_id = hashlib.sha256((group_id + members_json).encode("utf-8")).hexdigest()[:24]
        canonical_tuple = None
        if decision == "unify" and canonical_pk not in (None, {}):
            if not isinstance(canonical_pk, dict) or any(field not in canonical_pk for field in ITEM_PRIMARY_KEY_FIELDS):
                self.send_json(400, {"error": "Selecciona una clave primaria completa para el artículo maestro."})
                return
            canonical_tuple = item_pk_tuple(canonical_pk)
            if canonical_tuple not in selected_keys:
                self.send_json(400, {"error": "El artículo maestro debe pertenecer al subgrupo confirmado."})
                return
        elif decision != "unify" and canonical_pk not in (None, {}):
            self.send_json(400, {"error": "Solo un subgrupo confirmado puede tener artículo maestro."})
            return
        now = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
        canonical_json = json.dumps(canonical_pk, ensure_ascii=False, sort_keys=True) if canonical_tuple else None
        reopened_all = False
        with sqlite3.connect(ITEM_REVIEW_DB) as db:
            db.execute("BEGIN IMMEDIATE")
            if decision == "pending":
                rows = db.execute("SELECT group_id,decision,members_json FROM item_group_decisions WHERE run_id=? AND source_group_id=? AND decision IN ('unify','skip')", (run_id, group_id)).fetchall()
                matches = [(existing_id, existing_decision, existing_json,
                            {item_pk_tuple(pk) for pk in json.loads(existing_json)})
                           for existing_id, existing_decision, existing_json in rows
                           if set(selected_keys).issubset({item_pk_tuple(pk) for pk in json.loads(existing_json)})]
                if len(matches) != 1:
                    self.send_json(409, {"error": "La selección debe pertenecer a un solo subgrupo confirmado u omitido."})
                    return
                existing_id, existing_decision, existing_json, existing_keys = matches[0]
                if set(selected_keys) != existing_keys:
                    if existing_decision != "unify":
                        self.send_json(409, {"error": "Solo se pueden devolver parcialmente artículos de un subgrupo confirmado."})
                        return
                    remaining_keys = existing_keys - set(selected_keys)
                    if len(remaining_keys) >= 2:
                        remaining_pks = [member_by_key[key] for key in sorted(remaining_keys)]
                        existing_canonical = db.execute(
                            "SELECT canonical_pk FROM item_group_decisions WHERE run_id=? AND group_id=?",
                            (run_id, existing_id),
                        ).fetchone()
                        retained_canonical = existing_canonical[0] if existing_canonical else None
                        if retained_canonical and item_pk_tuple(json.loads(retained_canonical)) not in remaining_keys:
                            retained_canonical = None
                        db.execute("UPDATE item_group_decisions SET members_json=?,canonical_pk=? WHERE run_id=? AND group_id=?",
                                   (json.dumps(remaining_pks, ensure_ascii=False, sort_keys=True), retained_canonical, run_id, existing_id))
                    else:
                        # A one-member remainder cannot stay confirmed as a same-item subgroup.
                        reopened_all = True
                        selected_keys = sorted(existing_keys)
                        member_pks = [member_by_key[key] for key in selected_keys]
                        members_json = json.dumps(member_pks, ensure_ascii=False, sort_keys=True)
                        subgroup_id = existing_id
                else:
                    subgroup_id = existing_id
            if decision != "pending":
                rows = db.execute("SELECT group_id,decision,members_json FROM item_group_decisions WHERE run_id=? AND source_group_id=?", (run_id, group_id)).fetchall()
                for existing_id, existing_decision, existing_json in rows:
                    existing_keys = {item_pk_tuple(pk) for pk in json.loads(existing_json)}
                    if existing_id == subgroup_id and existing_decision == "unify" and existing_keys == set(selected_keys):
                        continue
                    if existing_decision in ("unify", "skip") and existing_keys.intersection(selected_keys):
                        self.send_json(409, {"error": "Uno o más artículos ya tienen una decisión guardada. Reábrela primero."})
                        return
            db.execute("""INSERT INTO item_group_decisions
                (run_id,group_id,decision,canonical_pk,members_json,saved_by,saved_at,source_group_id) VALUES (?,?,?,?,?,?,?,?)
                ON CONFLICT(run_id,group_id) DO UPDATE SET decision=excluded.decision,
                canonical_pk=excluded.canonical_pk,members_json=excluded.members_json,
                saved_by=excluded.saved_by,saved_at=excluded.saved_at,source_group_id=excluded.source_group_id""",
                (run_id, subgroup_id, decision, canonical_json, members_json, username, now, group_id))
            db.execute("""INSERT INTO item_group_decision_audit
                (run_id,group_id,decision,canonical_pk,modified_by,modified_at,source_group_id,members_json) VALUES (?,?,?,?,?,?,?,?)""",
                (run_id, subgroup_id, decision, canonical_json, username, now, group_id, members_json))
        self.send_json(200, {"saved": True, "group_id": group_id, "decision_id": subgroup_id, "decision": decision,
                             "canonical_pk": canonical_pk, "saved_by": username, "saved_at": now,
                             "reopened_member_count": len(member_pks) if decision == "pending" else 0,
                             "reopened_all": reopened_all})

    def start_items_compare(self, tipo, threshold, auto_same_threshold):
        global items_compare_running
        try:
            row_count = item_type_row_count(tipo)
        except Exception:
            self.send_json(503, {"error": "No se pudo validar el tipo seleccionado en Oracle."})
            return
        if row_count == 0:
            self.send_json(404, {"error": "El tipo seleccionado ya no tiene registros."})
            return
        with items_jobs_lock:
            if items_compare_running or items_export_running or full_export_running:
                self.send_json(409, {"error": "Ya hay otra exportación o comparación en curso."})
                return
            job_id = uuid.uuid4().hex
            output_dir = ITEM_EXPORT_ROOT / f"compare_{job_id}"
            try:
                output_dir.mkdir(parents=True, exist_ok=False)
            except OSError:
                self.send_json(507, {"error": "No se pudo preparar el directorio temporal."})
                return
            items_compare_running = True
            items_compare_jobs[job_id] = {
                "id": job_id,
                "tipo": tipo,
                "threshold": threshold,
                "auto_same_threshold": auto_same_threshold,
                "created_by": request_context.username,
                "state": "queued",
                "rows_processed": 0,
                "expected_rows": row_count,
                "current_batch": 0,
                "results": [],
                "created_at": time.time(),
            }
            try:
                create_comparison_run(job_id, tipo, threshold, auto_same_threshold, request_context.username)
            except sqlite3.Error:
                items_compare_jobs.pop(job_id, None)
                items_compare_running = False
                shutil.rmtree(output_dir, ignore_errors=True)
                self.send_json(500, {"error": "No se pudo guardar la corrida en SQLite."})
                return
        worker = threading.Thread(
            target=run_items_compare,
            args=(job_id, tipo, threshold, row_count, output_dir, request_context.credentials),
            daemon=True,
        )
        worker.start()
        self.send_json(202, {
            "id": job_id,
            "tipo": tipo,
            "auto_same_threshold": auto_same_threshold,
            "created_by": request_context.username,
            "state": "queued",
            "expected_rows": row_count,
        })

    def save_items_compare_decision(self, job_id, payload):
        index = payload.get("index")
        decision = payload.get("decision")
        keep = payload.get("keep")
        if not isinstance(index, int) or isinstance(index, bool):
            self.send_json(400, {"error": "El candidato seleccionado no es válido."})
            return
        if comparison_run_state(job_id) != "completed":
            self.send_json(404, {"error": "La corrida no está disponible para registrar decisiones."})
            return
        with items_jobs_lock:
            job = items_compare_jobs.get(job_id)
            pair = job["results"][index] if job and job.get("state") == "completed" and 0 <= index < len(job["results"]) else None
        if pair is None:
            pair = load_run_pair(job_id, index)
        if pair is None:
            self.send_json(404, {"error": "No se encontró ese candidato en la comparación guardada."})
            return
        refresh_pair_assignments(job_id, [pair])
        assigned_to = (pair.get("assignment") or {}).get("assigned_to", "")
        if not assigned_to or assigned_to.casefold() != request_context.username.casefold():
            self.send_json(403, {"error": "Este par no está asignado a tu usuario Oracle."})
            return
        try:
            saved = save_item_decision(pair, decision, keep, job_id)
        except ValueError as exc:
            self.send_json(400, {"error": str(exc)})
            return
        pair.update(saved)
        pair.setdefault("audit_history", []).append({
            "decision": saved["decision"], "keep": saved["keep"],
            "canonical_pk": saved["canonical_pk"], "modified_by": saved["saved_by"],
            "modified_at": saved["saved_at"],
        })
        self.send_json(200, {"saved": True, **saved})

    def download_items_excel(self, job_id):
        with items_jobs_lock:
            job = items_jobs.get(job_id)
            if job is None or job.get("state") != "completed":
                self.send_json(404, {"error": "El Excel no está listo para descargar."})
                return
            file_path = (ROOT / job["output_file"]).resolve()
            filename = job["filename"]
        if ITEM_EXPORT_ROOT.resolve() not in file_path.parents or not file_path.is_file():
            self.send_json(404, {"error": "No se encontró el Excel generado."})
            return
        try:
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Length", str(file_path.stat().st_size))
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.end_headers()
            with file_path.open("rb") as excel_file:
                shutil.copyfileobj(excel_file, self.wfile, length=1024 * 1024)
        except (BrokenPipeError, ConnectionResetError):
            pass

    def start_items_export(self, tipo):
        global items_export_running
        try:
            row_count = item_type_row_count(tipo)
        except Exception:
            self.send_json(503, {"error": "No se pudo validar el tipo seleccionado en Oracle."})
            return
        if row_count == 0:
            self.send_json(404, {"error": "El tipo seleccionado ya no tiene registros."})
            return
        with items_jobs_lock:
            if items_export_running or items_compare_running or full_export_running:
                self.send_json(409, {"error": "Ya hay una exportación de SIS.ITEMS en curso."})
                return
            job_id = uuid.uuid4().hex
            output_dir = ITEM_EXPORT_ROOT / job_id
            try:
                output_dir.mkdir(parents=True, exist_ok=False)
            except OSError:
                self.send_json(507, {"error": "No se pudo preparar el directorio de salida."})
                return
            safe_type = re.sub(r"[^A-Za-z0-9_-]", "_", tipo)[:30] or "tipo"
            filename = f"SIS_ITEMS_TIPO_{safe_type}.xlsx"
            items_export_running = True
            items_jobs[job_id] = {
                "id": job_id,
                "tipo": tipo,
                "state": "queued",
                "rows_processed": 0,
                "expected_rows": row_count,
                "current_batch": 0,
                "filename": filename,
                "created_at": time.time(),
            }
        worker = threading.Thread(
            target=run_items_export,
            args=(job_id, tipo, row_count, output_dir, request_context.credentials),
            daemon=True,
        )
        worker.start()
        self.send_json(202, {
            "id": job_id,
            "tipo": tipo,
            "state": "queued",
            "expected_rows": row_count,
            "status_url": f"/api/v1/items/exports/{job_id}",
        })

    def download_parquet(self, token):
        cleanup_expired_downloads()
        with downloads_lock:
            record = downloads.pop(token, None)
        if record is None or not record["file"].is_file():
            self.send_json(404, {"error": "La descarga ya expiró o fue utilizada."})
            return
        try:
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.apache.parquet")
            self.send_header("Content-Length", str(record["file"].stat().st_size))
            self.send_header("Content-Disposition", f'attachment; filename="{record["filename"]}"')
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.end_headers()
            with record["file"].open("rb") as parquet_file:
                shutil.copyfileobj(parquet_file, self.wfile, length=1024 * 1024)
        except (BrokenPipeError, ConnectionResetError):
            pass
        finally:
            shutil.rmtree(record["directory"], ignore_errors=True)

    def full_export_status(self, job_id):
        with full_jobs_lock:
            job = full_jobs.get(job_id)
            if job is None:
                self.send_json(404, {"error": "No se encontró la exportación."})
                return
            response = {key: value for key, value in job.items() if key not in ("process", "started_at", "statistics")}
            response["estimated_rows"] = job["statistics"]["estimated_rows"]
        self.send_json(200, response)

    def download_full_export(self, job_id):
        with full_jobs_lock:
            job = full_jobs.get(job_id)
            if job is None or job.get("state") != "completed":
                self.send_json(404, {"error": "El archivo no está listo para descargar."})
                return
            file_path = (ROOT / job["output_file"]).resolve()
            filename = f"{job['schema']}_{job['table']}.parquet"
        if ROOT not in file_path.parents or not file_path.is_file():
            self.send_json(404, {"error": "No se encontró el archivo Parquet."})
            return
        try:
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.apache.parquet")
            self.send_header("Content-Length", str(file_path.stat().st_size))
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.end_headers()
            with file_path.open("rb") as parquet_file:
                shutil.copyfileobj(parquet_file, self.wfile, length=1024 * 1024)
        except (BrokenPipeError, ConnectionResetError):
            pass

    def start_full_export(self, schema, table):
        global full_export_running
        try:
            estimate = estimate_table(schema, table)
        except ValueError as exc:
            self.send_json(422, {"error": str(exc)})
            return
        except Exception:
            self.send_json(503, {"error": "No se pudieron consultar estadísticas Oracle para estimar el tamaño."})
            return

        required_bytes = int(estimate["estimated_raw_bytes"] * 1.5 + 512 * 1024 * 1024)
        if required_bytes > estimate["free_bytes"]:
            self.send_json(507, {
                "error": "El espacio libre no ofrece margen suficiente para una exportación completa.",
                **estimate,
                "required_bytes": required_bytes,
            })
            return

        with full_jobs_lock:
            if full_export_running or items_export_running or items_compare_running:
                self.send_json(409, {"error": "Ya hay una exportación completa en curso."})
                return
            job_id = uuid.uuid4().hex
            output_dir = ROOT / "output" / "exports" / f"{schema}_{table}_{job_id[:8]}"
            try:
                output_dir.mkdir(parents=True, exist_ok=False)
            except OSError:
                self.send_json(507, {"error": "No se pudo crear el directorio de exportación."})
                return
            full_export_running = True
            full_jobs[job_id] = {
                "id": job_id, "schema": schema, "table": table,
                "state": "queued", "rows_processed": 0, "current_batch": 0,
                "bytes_written": 0, "elapsed_ms": 0,
                "output_directory": str(output_dir.relative_to(ROOT)),
                "statistics": estimate,
            }
        worker = threading.Thread(
            target=run_full_export,
            args=(job_id, schema, table, output_dir, estimate, request_context.credentials),
            daemon=True,
        )
        worker.start()
        self.send_json(202, {
            "id": job_id,
            "state": "queued",
            "estimated_rows": estimate["estimated_rows"],
            "estimated_raw_bytes": estimate["estimated_raw_bytes"],
            "free_bytes": estimate["free_bytes"],
            "required_bytes": required_bytes,
            "output_directory": str(output_dir.relative_to(ROOT)),
            "statistics_last_analyzed": estimate["last_analyzed"],
        })

    def catalog_schemas(self):
        try:
            with oracle_connection() as connection:
                with connection.cursor() as cursor:
                    cursor.execute("SELECT DISTINCT owner FROM all_tables ORDER BY owner")
                    items = [{"name": row[0]} for row in cursor]
            self.send_json(200, {"items": items})
        except Exception:
            self.send_json(503, {"error": "No se pudo consultar el catálogo Oracle local."})

    def catalog_tables(self, schema):
        try:
            with oracle_connection() as connection:
                with connection.cursor() as cursor:
                    cursor.execute("SELECT table_name FROM all_tables WHERE owner = :owner ORDER BY table_name", owner=schema)
                    items = [{"schema": schema, "name": row[0]} for row in cursor]
            self.send_json(200, {"items": items})
        except Exception:
            self.send_json(503, {"error": "No se pudieron consultar las tablas visibles de ese esquema."})

    def catalog_columns(self, schema, table):
        try:
            with oracle_connection() as connection:
                with connection.cursor() as cursor:
                    cursor.execute(
                        "SELECT column_name, data_type, nullable, data_precision, data_scale, data_length "
                        "FROM all_tab_columns WHERE owner = :owner AND table_name = :table_name ORDER BY column_id",
                        owner=schema, table_name=table,
                    )
                    items = [{
                        "name": row[0], "type": row[1], "nullable": row[2] == "Y",
                        "precision": row[3], "scale": row[4], "byte_length": row[5],
                    } for row in cursor]
            if not items:
                self.send_json(404, {"error": "La tabla no está visible para esta conexión."})
                return
            self.send_json(200, {"schema": schema, "name": table, "columns": items})
        except Exception:
            self.send_json(503, {"error": "No se pudieron consultar los metadatos de la tabla."})

    def do_POST(self):
        cleanup_expired_downloads()
        if not self.valid_local_request():
            self.send_json(403, {"error": "Solicitud local no válida."})
            return
        route = urlsplit(self.path).path
        decision_match = re.fullmatch(r"/api/v1/items/compare/([0-9a-f]{32})/decisions", route)
        group_decision_match = re.fullmatch(r"/api/v1/items/compare/([0-9a-f]{32})/groups/decisions", route)
        assignment_import_match = re.fullmatch(r"/api/v1/items/compare/([0-9a-f]{32})/assignments/import", route)
        allowed_routes = ("/api/v1/login", "/api/v1/logout", "/api/v1/items/compare")
        if route not in allowed_routes and decision_match is None and group_decision_match is None and assignment_import_match is None:
            self.send_json(404, {"error": "Ruta no encontrada."})
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            limit = 20 * 1024 * 1024 if assignment_import_match else 131072
            if length < 0 or length > limit or (assignment_import_match and length == 0):
                self.send_json(413, {"error": "Solicitud demasiado grande."})
                return
            body = self.rfile.read(length)
            payload = {} if assignment_import_match or not body else json.loads(body)
        except (ValueError, json.JSONDecodeError):
            self.send_json(400, {"error": "Solicitud JSON no válida."})
            return
        if route == "/api/v1/login":
            username, password = payload.get("username"), payload.get("password")
            if not isinstance(username, str) or not isinstance(password, str) or len(username) > 128 or len(password) > 512:
                self.send_json(400, {"error": "Usuario o contraseña no válidos."})
                return
            if not validate_oracle_login(username, password):
                self.send_json(401, {"error": "No se pudo autenticar con Oracle. Verifica tus credenciales."})
                return
            token = secrets.token_urlsafe(32)
            with sessions_lock:
                sessions[token] = {"username": username, "credentials": (username, password),
                                   "expires_at": time.time() + SESSION_TTL}
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Cache-Control", "no-store")
            self.send_header("Set-Cookie", f"zig_session={token}; HttpOnly; SameSite=Strict; Path=/; Max-Age={SESSION_TTL}")
            self.end_headers()
            self.wfile.write(b'{"authenticated":true}')
            return
        token, session = session_for(self)
        if route != "/api/v1/logout" and not session:
            self.send_json(401, {"error": "Inicia sesión con tu usuario Oracle."})
            return
        if session:
            request_context.credentials = session["credentials"]
            request_context.username = session["username"]
        if route == "/api/v1/logout":
            if token:
                with sessions_lock:
                    sessions.pop(token, None)
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Set-Cookie", "zig_session=; HttpOnly; SameSite=Strict; Path=/; Max-Age=0")
            self.end_headers()
            self.wfile.write(b'{"authenticated":false}')
            return
        if decision_match:
            self.save_items_compare_decision(decision_match.group(1), payload)
            return
        if group_decision_match:
            try:
                self.save_items_compare_group_decision(group_decision_match.group(1), payload)
            except sqlite3.Error:
                self.send_json(503, {"error": "No se pudo guardar la decisión del grupo en SQLite."})
            return
        if assignment_import_match:
            if request_context.username.casefold() != COMPARISON_RUNNER_USER:
                self.send_json(403, {"error": "Solo la cuenta responsable puede cargar asignaciones."})
                return
            try:
                result = import_assignment_workbook(assignment_import_match.group(1), body, request_context.username)
            except ValueError as error:
                self.send_json(400, {"error": str(error)})
                return
            except sqlite3.Error:
                self.send_json(503, {"error": "No se pudo guardar la asignación en SQLite."})
                return
            self.send_json(200, result)
            return
        if route == "/api/v1/items/compare":
            if request_context.username.casefold() != COMPARISON_RUNNER_USER:
                self.send_json(403, {"error": "Solo la cuenta responsable puede crear corridas de comparación."})
                return
            tipo = payload.get("tipo")
            threshold = payload.get("threshold", 85)
            auto_same_threshold = payload.get("auto_same_threshold", 94)
            if not isinstance(tipo, str) or not tipo or len(tipo) > 30 or "\x00" in tipo:
                self.send_json(400, {"error": "Selecciona un tipo SIS.ITEMS válido."})
                return
            if not isinstance(threshold, int) or isinstance(threshold, bool) or not 70 <= threshold <= 98:
                self.send_json(400, {"error": "El umbral debe estar entre 70 y 98."})
                return
            if not isinstance(auto_same_threshold, int) or isinstance(auto_same_threshold, bool) or not threshold <= auto_same_threshold <= 100:
                self.send_json(400, {"error": "El umbral para sugerir mismo artículo debe ser igual o mayor a similitud mínima y no superar 100."})
                return
            self.start_items_compare(tipo, threshold, auto_same_threshold)
            return
        if route == "/api/v1/items/exports":
            tipo = payload.get("tipo")
            if not isinstance(tipo, str) or not tipo or len(tipo) > 30 or "\x00" in tipo:
                self.send_json(400, {"error": "Selecciona un tipo SIS.ITEMS válido."})
                return
            self.start_items_export(tipo)
            return
        schema = payload.get("schema", "TABLEAU")
        table = payload.get("table", "ANALISISLABORATORIO")
        if not isinstance(schema, str) or not isinstance(table, str) or not valid_identifier(schema) or not valid_identifier(table):
            self.send_json(400, {"error": "Selecciona un esquema y una tabla válidos."})
            return
        if route == "/api/v1/full-exports":
            self.start_full_export(schema, table)
            return
        started = time.monotonic()
        cleanup_expired_downloads()
        output_dir = Path(tempfile.mkdtemp(prefix="zig-proyecto-parquet-download-"))
        try:
            result = subprocess.run(
                ["bash", str(PROBE)], cwd=ROOT, capture_output=True, text=True,
                timeout=120, check=False,
                env={**os.environ, "ORACLE_SCHEMA": schema, "ORACLE_TABLE": table,
                     "PARQUET_OUTPUT_DIR": str(output_dir)},
            )
        except subprocess.TimeoutExpired:
            shutil.rmtree(output_dir, ignore_errors=True)
            self.send_json(504, {"error": "La prueba excedió el tiempo permitido."})
            return
        output = (result.stdout + result.stderr)[-MAX_OUTPUT:]
        exported = EXPORT_RE.search(output)
        verified = VERIFY_RE.search(output)
        if result.returncode != 0 or not exported or not verified:
            shutil.rmtree(output_dir, ignore_errors=True)
            self.send_json(502, {"error": "Falló la conexión o la validación Oracle/Parquet. Revisa .env y los requisitos locales."})
            return
        if exported.group(1) != schema or exported.group(2) != table:
            shutil.rmtree(output_dir, ignore_errors=True)
            self.send_json(502, {"error": "El resultado no coincide con la tabla seleccionada."})
            return
        parquet_path = output_dir / "part-00001.parquet"
        if not parquet_path.is_file():
            shutil.rmtree(output_dir, ignore_errors=True)
            self.send_json(502, {"error": "No se encontró el Parquet generado."})
            return
        token = secrets.token_urlsafe(32)
        with downloads_lock:
            downloads[token] = {
                "directory": output_dir,
                "file": parquet_path,
                "filename": f"{schema}_{table}_muestra.parquet",
                "expires_at": time.time() + DOWNLOAD_TTL,
            }
        self.send_json(200, {
            "state": "completed",
            "schema": exported.group(1),
            "table": exported.group(2),
            "rows": int(exported.group(3)),
            "columns": int(verified.group(2)),
            "bytes": int(exported.group(5)),
            "elapsed_ms": round(float(exported.group(7)), 3),
            "verification": "PyArrow abrió el Parquet correctamente",
            "download_url": f"/api/v1/downloads/{token}",
            "download_expires_in_seconds": DOWNLOAD_TTL,
            "temporary_output_removed": False,
            "wall_ms": round((time.monotonic() - started) * 1000, 1),
        })

    def do_DELETE(self):
        if not self.valid_local_request():
            self.send_json(403, {"error": "Solicitud local no válida."})
            return
        match = re.fullmatch(r"/api/v1/items/compare/runs/([0-9a-f]{32})", urlsplit(self.path).path)
        if not match:
            self.send_json(404, {"error": "Ruta no encontrada."})
            return
        token, session = session_for(self)
        if not session:
            self.send_json(401, {"error": "Inicia sesión con tu usuario Oracle."})
            return
        if session["username"].casefold() != COMPARISON_RUNNER_USER:
            self.send_json(403, {"error": "Solo Christian Ruiz puede retirar corridas del historial."})
            return
        try:
            if not delete_comparison_run(match.group(1), session["username"]):
                self.send_json(404, {"error": "La corrida no existe, no terminó o ya fue retirada."})
                return
        except sqlite3.Error:
            self.send_json(500, {"error": "No se pudo registrar el retiro de la corrida."})
            return
        with items_jobs_lock:
            items_compare_jobs.pop(match.group(1), None)
        self.send_json(200, {"deleted": True, "audit_preserved": True})


if __name__ == "__main__":
    cleanup_expired_downloads()
    print(f"Prueba web Oracle → Parquet: http://{HOST}:{PORT}/probe", flush=True)
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
